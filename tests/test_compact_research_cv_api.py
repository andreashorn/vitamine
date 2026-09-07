import os
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from vitamine.app import app, export_format_catalog
from vitamine.paths import create_blank_database


FORMAT_ID = "vitamine.compact-research-cv"


class CompactResearchCvApiTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        root = Path(self.directory.name)
        self.database = root / "synthetic.vitamine"
        self.output = root / "output"
        create_blank_database(self.database)
        with sqlite3.connect(self.database) as con:
            con.execute(
                "UPDATE person SET full_name='Alex Example', display_name='Alex Example', "
                "position_title='Research Fellow', work_email='alex@example.invalid' WHERE id=1"
            )
            document_id = con.execute("SELECT id FROM documents LIMIT 1").fetchone()[0]
            con.execute(
                "INSERT INTO cv_entries "
                "(document_id, section_key, start_date, end_date, title, organization, raw_text, include_short) "
                "VALUES (?, 'education', '2014', '2018', 'PhD', 'Example University', 'PhD', 1)",
                (document_id,),
            )
            con.execute(
                "INSERT INTO publications "
                "(document_id, category, authors, title, venue, year, raw_citation, include_short) "
                "VALUES (?, 'peer_reviewed', 'Example A', 'A synthetic research result', "
                "'Example Journal', '2026', 'Example A. A synthetic research result. 2026.', 1)",
                (document_id,),
            )
        self.preferences = {"unrelated_preference": "preserved"}
        patches = (
            patch.dict(os.environ, {"VITAMINE_DB": str(self.database), "VITAMINE_OUTPUT": str(self.output)}),
            patch("vitamine.app.OUTPUT", self.output),
            patch("vitamine.paths.OUTPUT", self.output),
            patch("vitamine.app.read_preferences", side_effect=lambda: dict(self.preferences)),
            patch("vitamine.app.write_preferences", side_effect=self.write_preferences),
        )
        for patcher in patches:
            patcher.start()
            self.addCleanup(patcher.stop)
        self.client = TestClient(app)
        self.addCleanup(self.client.close)

    def write_preferences(self, payload):
        self.preferences.clear()
        self.preferences.update(payload)

    def listed_format(self):
        response = self.client.get("/api/export-formats")
        self.assertEqual(response.status_code, 200, response.text)
        return next(item for item in response.json()["formats"] if item["id"] == FORMAT_ID)

    def test_catalog_and_install_lifecycle(self):
        item = next(item for item in export_format_catalog() if item["id"] == FORMAT_ID)
        self.assertEqual(item["name"], "Compact Research CV")
        self.assertEqual(item["exporter"], "compact_research")
        self.assertEqual(item["content_profile"], "short")
        self.assertEqual(item["languages"], ["en"])
        self.assertEqual(item["quality"]["key"], "ready")
        self.assertFalse(item["preinstalled"])
        self.assertNotIn("page_limit", item)
        self.assertFalse(self.listed_format()["installed"])

        blocked = self.client.post(f"/api/actions/export/{FORMAT_ID}")
        self.assertEqual(blocked.status_code, 409)
        installed = self.client.post(f"/api/export-formats/{FORMAT_ID}/install")
        self.assertEqual(installed.status_code, 200, installed.text)
        self.assertTrue(self.listed_format()["installed"])
        self.client.post(f"/api/export-formats/{FORMAT_ID}/install")
        self.assertEqual(self.preferences["installed_export_format_ids"].count(FORMAT_ID), 1)

        removed = self.client.delete(f"/api/export-formats/{FORMAT_ID}/install")
        self.assertEqual(removed.status_code, 200, removed.text)
        self.assertFalse(self.listed_format()["installed"])
        self.assertEqual(self.preferences["unrelated_preference"], "preserved")
        blocked = self.client.post(f"/api/actions/export/{FORMAT_ID}")
        self.assertEqual(blocked.status_code, 409)

    def test_installed_format_exports_current_database_without_ai_or_canonical_build(self):
        self.client.post(f"/api/export-formats/{FORMAT_ID}/install")
        with (
            patch("vitamine.app.llm_json", side_effect=AssertionError("Unexpected AI call")),
            patch("vitamine.app.run_export_quality_audit", side_effect=AssertionError("Unexpected AI audit")),
            patch("vitamine.app.build_short_action", side_effect=AssertionError("Unexpected canonical build")),
            patch("vitamine.app.run_script", side_effect=AssertionError("Unexpected subprocess")),
        ):
            response = self.client.post(f"/api/actions/export/{FORMAT_ID}?lang=en")
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertTrue(payload["ok"])
        self.assertEqual(payload["template_id"], FORMAT_ID)
        self.assertEqual(payload["language"], "en")
        self.assertEqual(payload["docx_path"], "output/compact_research_cv_en.docx")
        self.assertIsInstance(payload["template_render"], dict)
        document = Document(self.output / "compact_research_cv_en.docx")
        text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        self.assertIn("Alex Example", text)
        self.assertIn("Example University", text)
        self.assertIn("A synthetic research result", text)
        self.assertNotIn("{{", text)

    def test_unsupported_languages_do_not_render(self):
        self.client.post(f"/api/export-formats/{FORMAT_ID}/install")
        with patch("vitamine.app.render_compact_research_cv") as render:
            for language in ("secondary", "de", "fr"):
                with self.subTest(language=language):
                    response = self.client.post(f"/api/actions/export/{FORMAT_ID}?lang={language}")
                    self.assertEqual(response.status_code, 422, response.text)
            render.assert_not_called()

    def test_render_failure_does_not_expose_document_data(self):
        self.client.post(f"/api/export-formats/{FORMAT_ID}/install")
        with patch("vitamine.app.render_compact_research_cv", side_effect=ValueError("private source detail")):
            response = self.client.post(f"/api/actions/export/{FORMAT_ID}")
        self.assertEqual(response.status_code, 500)
        self.assertFalse(response.json()["ok"])
        self.assertNotIn("private source detail", response.text)


if __name__ == "__main__":
    unittest.main()
