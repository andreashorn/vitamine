import unittest

from docx import Document

from vitamine.scripts.build_biosketch import set_run_font as set_biosketch_font
from vitamine.scripts.build_long_cv import set_run_font as set_long_font
from vitamine.scripts.build_short_cv import set_paragraph_font as set_short_font
from vitamine.scripts.build_ultrashort_tabular_cv import set_run_font as set_ultrashort_font


class ExportFontFamilyTests(unittest.TestCase):
    def test_cv_exporters_use_helvetica_but_biosketch_retains_arial(self):
        document = Document()

        long_run = document.add_paragraph().add_run("Long")
        set_long_font(long_run)
        short_paragraph = document.add_paragraph("Short")
        set_short_font(short_paragraph)
        ultrashort_run = document.add_paragraph().add_run("Ultrashort")
        set_ultrashort_font(ultrashort_run)
        biosketch_run = document.add_paragraph().add_run("Biosketch")
        set_biosketch_font(biosketch_run)

        self.assertEqual(long_run.font.name, "Helvetica")
        self.assertTrue(all(run.font.name == "Helvetica" for run in short_paragraph.runs))
        self.assertEqual(ultrashort_run.font.name, "Helvetica")
        self.assertEqual(biosketch_run.font.name, "Arial")


if __name__ == "__main__":
    unittest.main()
