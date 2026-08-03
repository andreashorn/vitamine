import io
import os
import sqlite3
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.shared import Inches, Pt
from fastapi.testclient import TestClient

from vitamine.app import app
from vitamine.custom_docx_templates import (
    analyze_and_skeletonize,
    canonical_content,
    deterministic_profile,
    render_template,
)
from vitamine.paths import create_blank_database


def document_bytes(*, biosketch: bool = False) -> bytes:
    document = Document()
    document.sections[0].left_margin = Inches(0.57)
    document.add_paragraph("Jane Example, PhD")
    document.add_paragraph("jane@example.org")
    if biosketch:
        document.add_heading("Biographical Sketch", level=1)
        document.add_heading("A. Personal Statement", level=1)
        document.add_paragraph("A private source narrative.")
        document.add_heading("C. Contributions to Science", level=1)
        document.add_paragraph("1. A private contribution.")
    else:
        document.add_heading("Education", level=1)
        table = document.add_table(rows=0, cols=2)
        cells = table.add_row().cells
        cells[0].text = "Years"
        cells[1].text = "Qualification"
        cells = table.add_row().cells
        cells[0].text = "2010-2014"
        cells[1].text = "PhD, Old University"
        document.add_heading("Selected Publications", level=1)
        document.add_paragraph("Example J. An old paper. 2024.")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def oxford_like_document_bytes() -> bytes:
    document = Document()
    document.sections[0].left_margin = Inches(0.57)
    document.styles["Heading 1"].paragraph_format.space_before = None
    document.add_paragraph("LAURA SOURCE").runs[0].bold = True
    document.add_paragraph(
        "Department of Example Sciences, Source University, Example Road, Oxford, "
        "OX1 6AY, laura.source@example.invalid"
    )
    document.add_heading("RESEARCH EXPERIENCE", level=1)
    document.add_heading("Postdoctoral Research Scientist, University of Oxford (2022-2025)", level=2)
    document.add_paragraph("A source-only research responsibility.", style="List Bullet")
    document.add_heading("RELEVANT RESEARCH SKILLS", level=1)
    document.add_paragraph("A source-only laboratory skill.", style="List Bullet")
    document.add_heading("AWARDS", level=1)
    document.add_paragraph("A source-only award")
    document.add_heading("EDUCATION", level=1)
    education = document.add_paragraph("University of Nottingham (2018-2022)")
    education.runs[0].bold = True
    document.add_paragraph("A source-only degree")
    document.add_heading("GENERAL SKILLS & COURSES", level=1)
    document.add_paragraph("A source-only unusual general skill")
    document.add_heading("TEACHING EXPERIENCE", level=1)
    document.add_heading("Mentor/Supervisor, University of Oxford (2022-2025)", level=2)
    document.add_paragraph("A source-only teaching duty.", style="List Bullet")
    document.add_heading("ADDITIONAL RELEVANT EXPERIENCE", level=1)
    document.add_heading("Assistant Information Officer (2015-2016)", level=2)
    document.add_paragraph("A source-only professional duty.", style="List Bullet")
    document.add_heading("MEMBERSHIP OF PROFESSIONAL SOCIETIES", level=1)
    document.add_paragraph("Source Society")
    document.add_heading("REFEREES", level=1)
    document.add_paragraph("Prof Source Referee")
    document.add_heading("PUBLICATIONS", level=1)
    document.add_paragraph("A source-only publication")
    document.add_heading("CONFERENCE PAPERS", level=2)
    document.add_paragraph("A source-only conference paper")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def semantic_canonical(path: Path) -> None:
    document = Document()
    document.add_heading("Academic Appointments", level=1)
    table = document.add_table(rows=0, cols=2)
    cells = table.add_row().cells
    cells[0].text = "2024-present"
    cells[1].text = "Professor of Current Science, Current University"
    document.add_heading("Honors and Awards", level=1)
    table = document.add_table(rows=0, cols=2)
    cells = table.add_row().cells
    cells[0].text = "2025"
    cells[1].text = "Current Research Prize"
    document.add_heading("Education", level=1)
    table = document.add_table(rows=0, cols=2)
    cells = table.add_row().cells
    cells[0].text = "2010-2014"
    cells[1].text = "PhD, Current University"
    document.add_heading("Selected Publications", level=1)
    publications = document.add_table(rows=0, cols=2)
    for index, title in enumerate(("First current paper", "Second current paper"), 1):
        cells = publications.add_row().cells
        cells[0].text = f"{index}."
        cells[1].text = title
    document.add_heading("Invited Presentations", level=1)
    table = document.add_table(rows=0, cols=2)
    cells = table.add_row().cells
    cells[0].text = "2026"
    cells[1].text = "A current invited talk that the source form does not request"
    document.save(path)


def memory_database() -> sqlite3.Connection:
    con = sqlite3.connect(":memory:")
    con.row_factory = sqlite3.Row
    con.execute(
        """
        CREATE TABLE person (
          id INTEGER PRIMARY KEY, full_name TEXT, display_name TEXT, degrees TEXT,
          position_title TEXT, office_address TEXT, home_address TEXT, work_phone TEXT,
          work_email TEXT, place_of_birth TEXT, era_commons TEXT, orcid_id TEXT
        )
        """
    )
    con.execute(
        "INSERT INTO person (id, full_name, work_email) VALUES (1, 'Jane Example', 'jane@example.org')"
    )
    return con


class CustomDocxTemplateUnitTests(unittest.TestCase):
    def test_template_analysis_uses_its_task_model_without_changing_the_default(self):
        observed = {}

        def fake_llm(_prompt, _schema, settings):
            observed.update(settings)
            return None, "test fallback"

        base_settings = {
            "provider": "openai",
            "api_model": "gpt-5.4-nano",
            "task_models": {"custom_template_analysis": "gpt-5.6-luna"},
        }
        analyze_and_skeletonize(
            document_bytes(),
            "Task-routed layout",
            memory_database(),
            llm_json=fake_llm,
            settings=base_settings,
        )
        self.assertEqual(observed["api_model"], "gpt-5.6-luna")
        self.assertEqual(base_settings["api_model"], "gpt-5.4-nano")

    def test_biosketch_is_classified_without_an_llm(self):
        con = memory_database()
        _skeleton, blueprint = analyze_and_skeletonize(
            document_bytes(biosketch=True),
            "Biosketch layout",
            con,
            settings={"provider": "none"},
        )
        self.assertEqual(blueprint["content_profile"], "biosketch")
        self.assertIn("personal_statement", blueprint["mapped_sections"])
        self.assertIn("contributions", blueprint["mapped_sections"])

    def test_roundtrip_preserves_page_geometry_and_replaces_private_content(self):
        con = memory_database()
        skeleton, blueprint = analyze_and_skeletonize(
            document_bytes(),
            "Faculty layout",
            con,
            settings={"provider": "none"},
        )
        skeleton_document = Document(io.BytesIO(skeleton))
        skeleton_text = "\n".join(
            [paragraph.text for paragraph in skeleton_document.paragraphs]
            + [cell.text for table in skeleton_document.tables for row in table.rows for cell in row.cells]
        )
        self.assertNotIn("Jane Example", skeleton_text)
        self.assertNotIn("Old University", skeleton_text)
        self.assertIn("VITAMINE_SECTION", skeleton_text)

        with tempfile.TemporaryDirectory() as folder:
            canonical = Document()
            canonical.add_heading("Education", level=1)
            table = canonical.add_table(rows=0, cols=2)
            cells = table.add_row().cells
            cells[0].text = "Years"
            cells[1].text = "Qualification"
            cells = table.add_row().cells
            cells[0].text = "2012-2016"
            cells[1].text = "MD, Current University"
            canonical.add_heading("Selected Publications", level=1)
            canonical.add_paragraph("Current A. An updated paper. 2026.")
            canonical_path = Path(folder) / "canonical.docx"
            canonical.save(canonical_path)
            output = Path(folder) / "rendered.docx"
            report = render_template(skeleton, blueprint, canonical_path, output, con)

            rendered = Document(output)
            rendered_text = "\n".join(
                [paragraph.text for paragraph in rendered.paragraphs]
                + [cell.text for table in rendered.tables for row in table.rows for cell in row.cells]
            )
            self.assertIn("Jane Example", rendered_text)
            self.assertIn("Current University", rendered_text)
            self.assertIn("updated paper", rendered_text)
            self.assertNotIn("Old University", rendered_text)
            self.assertAlmostEqual(rendered.sections[0].left_margin.inches, 0.57, places=2)
            self.assertEqual(report["rendered_sections"], ["education", "publications"])

    def test_curated_four_page_shape_is_short_not_long(self):
        text = " ".join(["research"] * 925)
        profile, _reason = deterministic_profile(text, page_count=4, heading_count=23)
        self.assertEqual(profile, "short")

    def test_publications_in_one_canonical_table_cell_remain_separate_items(self):
        document = Document()
        document.add_heading("Peer-Reviewed Publications", level=1)
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.paragraphs[0].text = "First canonical citation"
        cell.add_paragraph("Second canonical citation")
        self.assertEqual(
            canonical_content(document)["publications"],
            [["First canonical citation"], ["Second canonical citation"]],
        )

    def test_semantic_roundtrip_replaces_identity_and_isolates_source_sections(self):
        con = memory_database()
        con.execute(
            "UPDATE person SET display_name='Jane Current', office_address='Current Institute, Berlin', "
            "work_email='jane.current@example.org' WHERE id=1"
        )
        def overbroad_llm(_prompt, _schema, _settings):
            return {
                "content_profile": "short",
                "confidence": "high",
                "heading_mappings": [
                    {"source_text": "GENERAL SKILLS & COURSES", "section_key": "research_skills"}
                ],
            }, None

        skeleton, blueprint = analyze_and_skeletonize(
            oxford_like_document_bytes(),
            "Oxford-like layout",
            con,
            llm_json=overbroad_llm,
            settings={"provider": "openai", "api_model": "test-model"},
        )
        skeleton_document = Document(io.BytesIO(skeleton))
        skeleton_text = "\n".join(paragraph.text for paragraph in skeleton_document.paragraphs)
        self.assertNotIn("LAURA SOURCE", skeleton_text)
        self.assertNotIn("laura.source@example.invalid", skeleton_text)
        self.assertEqual(blueprint["schema_version"], 2)
        self.assertIn("research_experience", blueprint["mapped_sections"])
        self.assertIn("research_skills", blueprint["mapped_sections"])
        self.assertIn("conference_papers", blueprint["mapped_sections"])
        self.assertEqual(
            blueprint["manual_sections"],
            ["GENERAL SKILLS & COURSES", "REFEREES", "CONFERENCE PAPERS"],
        )

        with tempfile.TemporaryDirectory() as folder:
            canonical_path = Path(folder) / "canonical.docx"
            semantic_canonical(canonical_path)
            output = Path(folder) / "rendered.docx"
            render_template(skeleton, blueprint, canonical_path, output, con)
            rendered = Document(output)
            paragraphs = [paragraph.text.strip() for paragraph in rendered.paragraphs if paragraph.text.strip()]
            rendered_text = "\n".join(paragraphs)

            self.assertEqual(paragraphs[0], "Jane Current")
            self.assertIn("Current Institute, Berlin", paragraphs[1])
            self.assertIn("jane.current@example.org", paragraphs[1])
            self.assertNotIn("LAURA SOURCE", rendered_text)
            self.assertNotIn("Mentor/Supervisor, University of Oxford", rendered_text)
            self.assertIn("Professor of Current Science, Current University", rendered_text)
            self.assertIn("Current Research Prize", rendered_text)
            self.assertNotIn("current invited talk", rendered_text)
            self.assertIn("GENERAL SKILLS & COURSES", rendered_text)
            self.assertIn("RELEVANT RESEARCH SKILLS", rendered_text)
            self.assertIn("TEACHING EXPERIENCE", rendered_text)
            self.assertIn("ADDITIONAL RELEVANT EXPERIENCE", rendered_text)
            self.assertIn("MEMBERSHIP OF PROFESSIONAL SOCIETIES", rendered_text)
            self.assertIn("REFEREES", rendered_text)
            self.assertIn("CONFERENCE PAPERS", rendered_text)
            self.assertEqual(rendered_text.count("[Please fill this section manually.]"), 7)
            self.assertNotIn("source-only", rendered_text)
            publication_paragraphs = [text for text in paragraphs if "current paper" in text]
            self.assertEqual(publication_paragraphs, ["First current paper", "Second current paper"])
            self.assertLess(max(map(len, publication_paragraphs)), 100)
            additional_skills = next(
                paragraph for paragraph in rendered.paragraphs
                if paragraph.text.strip() == "GENERAL SKILLS & COURSES"
            )
            self.assertEqual(additional_skills.paragraph_format.space_before, Pt(6))

            with zipfile.ZipFile(output) as archive:
                self.assertFalse(any(name.startswith("word/header") for name in archive.namelist()))
                self.assertFalse(any(name.startswith("word/footer") for name in archive.namelist()))


class CustomDocxTemplateApiTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        root = Path(self.directory.name)
        self.database = root / "custom-template.vitamine"
        self.output = root / "output"
        create_blank_database(self.database)
        with sqlite3.connect(self.database) as con:
            con.execute(
                "UPDATE person SET full_name='Jane Example', display_name='Jane Example', work_email='jane@example.org' WHERE id=1"
            )
            document_id = con.execute("SELECT id FROM documents LIMIT 1").fetchone()[0]
            con.execute(
                """
                INSERT INTO cv_entries
                  (document_id, section_key, start_date, end_date, title, organization, raw_text,
                   include_long, include_short)
                VALUES (?, 'education', '2012', '2016', 'MD', 'Current University',
                        'MD, Current University', 1, 1)
                """,
                (document_id,),
            )
            con.execute(
                """
                INSERT INTO publications
                  (document_id, category, authors, title, venue, year, raw_citation,
                   include_short, include_ultrashort)
                VALUES (?, 'peer_reviewed', 'Example J', 'An updated paper', 'Science', '2026',
                        'Example J. An updated paper. Science. 2026.', 1, 1)
                """,
                (document_id,),
            )
        self.environment = patch.dict(
            os.environ,
            {"VITAMINE_DB": str(self.database), "VITAMINE_OUTPUT": str(self.output)},
        )
        self.environment.start()
        self.app_output = patch("vitamine.app.OUTPUT", self.output)
        self.paths_output = patch("vitamine.paths.OUTPUT", self.output)
        self.app_output.start()
        self.paths_output.start()
        self.llm = patch("vitamine.app.llm_json", return_value=(None, "disabled in test"))
        self.llm.start()
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        self.llm.stop()
        self.paths_output.stop()
        self.app_output.stop()
        self.environment.stop()
        self.directory.cleanup()

    def test_upload_rename_list_and_delete_private_word_template(self):
        created = self.client.post(
            "/api/export-templates",
            data={"name": "My Faculty Form"},
            files={
                "file": (
                    "faculty.docx",
                    document_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(created.status_code, 200, created.text)
        custom = created.json()["format"]
        self.assertTrue(custom["custom_template"])
        self.assertTrue(custom["installed"])
        self.assertIn(custom["content_profile"], {"short", "one_page"})

        listed = self.client.get("/api/export-formats")
        self.assertEqual(listed.status_code, 200, listed.text)
        self.assertEqual(listed.json()["formats"][0]["id"], custom["id"])

        renamed = self.client.put(
            f"/api/export-templates/{custom['id']}",
            json={"name": "Promotion dossier form"},
        )
        self.assertEqual(renamed.status_code, 200, renamed.text)
        self.assertEqual(renamed.json()["format"]["name"], "Promotion dossier form")

        exported = self.client.post(f"/api/actions/export/{custom['id']}?lang=en")
        self.assertEqual(exported.status_code, 200, exported.text)
        export_path = self.output / Path(exported.json()["docx_path"]).name
        self.assertTrue(export_path.exists())
        exported_document = Document(export_path)
        export_text = "\n".join(
            [paragraph.text for paragraph in exported_document.paragraphs]
            + [cell.text for table in exported_document.tables for row in table.rows for cell in row.cells]
        )
        self.assertIn("Current University", export_text)
        self.assertIn("updated paper", export_text)
        self.assertNotIn("Old University", export_text)

        with sqlite3.connect(self.database) as con:
            stored = con.execute(
                "SELECT source_docx, source_sha256 FROM export_templates WHERE id=?",
                (custom["id"],),
            ).fetchone()
        self.assertIsNotNone(stored)
        self.assertEqual(len(stored[1]), 64)
        stored_document = Document(io.BytesIO(stored[0]))
        self.assertNotIn("Jane Example", "\n".join(paragraph.text for paragraph in stored_document.paragraphs))

        deleted = self.client.delete(f"/api/export-templates/{custom['id']}")
        self.assertEqual(deleted.status_code, 200, deleted.text)
        listed = self.client.get("/api/export-formats").json()["formats"]
        self.assertFalse(any(item["id"] == custom["id"] for item in listed))

    def test_export_surface_contains_second_docx_dropzone(self):
        page = self.client.get("/")
        self.assertIn('id="customTemplateDropzone"', page.text)
        self.assertIn('id="customTemplateName"', page.text)
        script = self.client.get("/static/app.js")
        self.assertIn("async function importCustomExportTemplate(file)", script.text)
        self.assertIn("renameCustomExportTemplate", script.text)

    def test_stale_remove_action_deletes_a_custom_template(self):
        created = self.client.post(
            "/api/export-templates",
            data={"name": "Legacy remove target"},
            files={
                "file": (
                    "legacy-remove.docx",
                    document_bytes(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        self.assertEqual(created.status_code, 200, created.text)
        template_id = created.json()["format"]["id"]

        removed = self.client.delete(f"/api/export-formats/{template_id}/install")
        self.assertEqual(removed.status_code, 200, removed.text)
        self.assertTrue(removed.json()["deleted"])
        self.assertFalse(
            any(item["id"] == template_id for item in self.client.get("/api/export-formats").json()["formats"])
        )


if __name__ == "__main__":
    unittest.main()
