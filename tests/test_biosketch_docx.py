import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches

from vitamine.scripts.build_biosketch import (
    add_contribution_paragraph,
    add_two_column_table,
)


class BiosketchDocxTests(unittest.TestCase):
    def test_table_geometry_is_fixed_in_twips_and_borderless(self):
        document = Document()
        widths = [Inches(1.1), Inches(6.1)]
        add_two_column_table(document, [("2025", "A scientific appointment")], left_width=1.1)

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "biosketch.docx"
            document.save(path)
            rendered = Document(path)
            table = rendered.tables[0]
            expected = [str(int(width.twips)) for width in widths]
            borders = table._tbl.tblPr.first_child_found_in("w:tblBorders")

            self.assertEqual(
                [column.get(qn("w:w")) for column in table._tbl.tblGrid],
                expected,
            )
            self.assertEqual(
                table._tbl.tblPr.first_child_found_in("w:tblW").get(qn("w:w")),
                str(sum(int(width.twips) for width in widths)),
            )
            self.assertEqual(
                [
                    cell._tc.get_or_add_tcPr().find(qn("w:tcW")).get(qn("w:w"))
                    for cell in table.rows[0].cells
                ],
                expected,
            )
            self.assertTrue(all(edge.get(qn("w:val")) == "nil" for edge in borders))

    def test_contribution_bolds_only_number_and_title(self):
        document = Document()
        add_contribution_paragraph(document, 2, "A scientific contribution", "Narrative remains readable.")
        runs = document.paragraphs[0].runs

        self.assertEqual([run.text for run in runs], ["2. A scientific contribution.", " Narrative remains readable."])
        self.assertEqual([run.bold for run in runs], [True, False])


if __name__ == "__main__":
    unittest.main()
