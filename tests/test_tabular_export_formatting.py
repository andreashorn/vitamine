import sqlite3
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn

from vitamine.scripts.build_ultrashort_tabular_cv import (
    add_publication_docx_text,
    clear_paragraph,
    remove_numbering,
)
from vitamine.scripts.export_utils import configure_researcher_name


class TabularExportFormattingTests(unittest.TestCase):
    def publication(self):
        con = sqlite3.connect(":memory:")
        con.row_factory = sqlite3.Row
        con.execute(
            """CREATE TABLE publications (
                 id INTEGER, authors TEXT, title TEXT, venue TEXT, year TEXT,
                 doi TEXT, url TEXT, short_citation TEXT, impact_factor REAL,
                 impact_factor_year TEXT)"""
        )
        con.execute(
            """INSERT INTO publications VALUES
               (1, 'First Author, Jane Example, Third Author, Fourth Author',
                'A useful result', 'NATURE COMMUNICATIONS', '2026',
                '10.1234/example', '', '', 12.0, '2025')"""
        )
        return con, con.execute("SELECT * FROM publications").fetchone()

    def test_empty_template_bullet_loses_numbering(self):
        doc = Document()
        paragraph = doc.add_paragraph("placeholder", style="List Bullet")
        numbering = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        numbering.get_or_add_numId().val = 1
        clear_paragraph(paragraph)
        remove_numbering(paragraph)
        self.assertEqual(paragraph.text, "")
        self.assertIsNone(paragraph._p.pPr.find(qn("w:numPr")))

    def test_default_citation_uses_arial_manual_number_styling_and_doi_link(self):
        con, publication = self.publication()
        self.addCleanup(con.close)
        doc = Document()
        paragraph = doc.add_paragraph()
        configure_researcher_name({"display_name": "Jane Example", "full_name": "Jane Example"})
        with patch("vitamine.scripts.build_ultrashort_tabular_cv.configured_citation_style", return_value="vitamine-long"):
            add_publication_docx_text(paragraph, 1, publication)
        self.assertTrue(paragraph.text.startswith("1.\t"))
        self.assertIn("et al. (2026).", paragraph.text)
        self.assertIsNone(paragraph._p.pPr.find(qn("w:numPr")))
        self.assertTrue(all(run.font.name == "Arial" for run in paragraph.runs))
        venue_runs = [run for run in paragraph.runs if "Nature Communications" in run.text]
        self.assertEqual(len(venue_runs), 1)
        self.assertTrue(venue_runs[0].italic)
        self.assertTrue(venue_runs[0].underline)
        researcher_runs = [run for run in paragraph.runs if "Example" in run.text]
        self.assertTrue(researcher_runs and researcher_runs[0].bold)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "citation.docx"
            doc.save(path)
            with zipfile.ZipFile(path) as package:
                document_xml = package.read("word/document.xml")
                relationships = package.read("word/_rels/document.xml.rels")
            self.assertIn(b"https://doi.org/10.1234/example", document_xml)
            self.assertIn(b"https://doi.org/10.1234/example", relationships)


if __name__ == "__main__":
    unittest.main()
