"""Synthetic-only mapping and archive privacy checks for Compact Research CV."""

from __future__ import annotations

import json
import posixpath
import sqlite3
import tempfile
import unittest
import zipfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

from vitamine.compact_research_cv import TEMPLATE_PATH, render_compact_research_cv
from vitamine.paths import create_blank_database


TODAY = date(2026, 9, 7)
CANARY = "UNMAPPED_PRIVATE_CANARY"


def add_entry(con: sqlite3.Connection, **values) -> int:
    fields = {"section_key": "education", "title": "Example entry", "raw_text": CANARY,
              "source_note": CANARY, "include_short": 1, **values}
    columns = ", ".join(fields)
    placeholders = ", ".join("?" for _ in fields)
    return con.execute(
        f"INSERT INTO cv_entries ({columns}) VALUES ({placeholders})", tuple(fields.values())
    ).lastrowid


def add_publication(con: sqlite3.Connection, **values) -> int:
    fields = {"category": "peer_reviewed", "authors": "Example J, Sample R",
              "title": "Example publication", "venue": "Example Journal", "year": "2026",
              "raw_citation": "", "include_short": 1, **values}
    columns = ", ".join(fields)
    placeholders = ", ".join("?" for _ in fields)
    return con.execute(
        f"INSERT INTO publications ({columns}) VALUES ({placeholders})", tuple(fields.values())
    ).lastrowid


def populate_example(con: sqlite3.Connection) -> None:
    """Populate a blank schema with invented data for visual rendering checks."""
    con.execute(
        "UPDATE person SET full_name=?, display_name=?, degrees=?, orcid_id=?, "
        "position_title=?, office_address=?, own_institution_name=?, raw_json=? WHERE id=1",
        ("Jane Alice Example", "Jane Example", "PhD, MSc", "0000-0000-0000-0000",
         "Professor of Experimental Science", "Example Institute, 12 Sample Road, Example City",
         "Example Institute", json.dumps({"unmapped": CANARY})),
    )
    for section, start, end, title, organization, description in (
        ("education", "2010", "2014", "Doctorate in experimental science", "Example University",
         "Thesis: reproducible methods for complex laboratory measurements"),
        ("education", "2007", "2010", "Master of Science", "Sample University",
         "Advanced training in computational methods and quantitative research"),
        ("academic_appointments", "2023", "Present", "Professor of experimental science", "Example Institute",
         "Directs a research group developing open methods for reproducible experiments"),
        ("academic_appointments", "2019", "2023", "Associate professor", "Sample University",
         "Led interdisciplinary projects and supervised doctoral researchers"),
        ("postdoctoral_training", "2014", "2019", "Postdoctoral fellow", "Sample Research Centre",
         "Developed measurement tools and contributed to international research projects"),
        ("hospital_appointments", "2018", "2019", "Visiting research scientist", "Example Research Hospital",
         "Coordinated a translational research collaboration"),
        ("honors", "2025", "", "Open Research Award", "Example Scientific Society",
         "For a reusable collection of open analysis methods"),
        ("honors", "2021", "", "Early Career Research Prize", "Sample Research Foundation",
         "Recognition of methodological contributions"),
        ("honors", "2015", "", "Visiting Fellowship", "Example Academy",
         "Competitive fellowship supporting international exchange"),
        ("editorial_activities", "2022", "Present", "Editorial board member", "Journal of Example Science",
         "Supports transparent reporting and reproducible publication practices"),
        ("committee_service", "2021", "Present", "Research methods committee", "Example Scientific Society",
         "Coordinates workshops and shared reporting recommendations"),
        ("grant_review", "2020", "Present", "Research funding reviewer", "Sample Research Foundation",
         "Reviews interdisciplinary investigator-led research proposals"),
    ):
        add_entry(con, section_key=section, start_date=start, end_date=end, title=title,
                  organization=organization, description=description)
    for index, title in enumerate(("Open measurement methods", "Shared experimental platforms"), 1):
        add_entry(con, section_key="funding", start_date="2025", end_date="2029", title=title,
                  organization="Example Research Council", role="Principal investigator",
                  amount=f"EUR {index * 250000:,}", grant_status="funded",
                  description="Developing and validating accessible research tools across partner laboratories")
    for index in range(1, 6):
        title = f"Reproducible experimental methods: synthetic study {index}"
        authors = "Example J, Sample R, Model T, Demo A, Fiction B"
        citation = (f"{authors}. {title}. Journal of Example Science. "
                    f"{2027 - index};{20 + index}(2):100–112.")
        add_publication(con, title=title, authors=authors, year=str(2027 - index),
                        venue="Journal of Example Science", raw_citation=citation,
                        doi=f"10.0000/example.{index}", short_selected_order=index)
    con.commit()


def document_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    return "\n".join(root.itertext())


class CompactResearchCvTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.database = Path(self.directory.name) / "synthetic.vitamine"
        self.output = Path(self.directory.name) / "compact.docx"
        create_blank_database(self.database)
        self.con = sqlite3.connect(self.database)
        self.con.row_factory = sqlite3.Row
        self.addCleanup(self.con.close)

    def render(self):
        result = render_compact_research_cv(self.con, self.output, today=TODAY)
        return document_text(self.output), result

    def test_identity_uses_mapped_fields_and_safe_fallbacks(self):
        populate_example(self.con)
        text, _ = self.render()
        for value in ("Jane Example", "PhD, MSc", "0000-0000-0000-0000",
                      "Professor of Experimental Science", "12 Sample Road"):
            self.assertIn(value, text)
        self.assertNotIn("Jane Alice Example", text)
        self.con.execute("UPDATE person SET display_name='', office_address='' WHERE id=1")
        text, _ = self.render()
        self.assertIn("Jane Alice Example", text)
        self.assertIn("Example Institute", text)

    def test_all_structured_entry_fields_survive_in_each_mapped_group(self):
        section_counts = {"education": 1, "positions": 5, "honors": 1, "activities": 5, "funding": 1}
        sections = (
            "education", "postdoctoral_training", "academic_appointments", "hospital_appointments",
            "professional_positions", "research_experience", "honors", "editorial_activities",
            "grant_review", "committee_service", "professional_societies", "community_service", "funding",
        )
        for section in sections:
            add_entry(self.con, section_key=section, start_date="2025", end_date="2029", grant_status="funded",
                      **{key: f"{section} {key}" for key in
                         ("title", "organization", "location", "role", "amount", "description")})
        text, result = self.render()
        for section in sections:
            for field in ("title", "organization", "location", "role", "amount", "description"):
                self.assertIn(f"{section} {field}", text)
        self.assertIn("2025–2029", text)
        for section, count in section_counts.items():
            self.assertEqual(result["section_counts"][section], count)

    def test_short_selections_and_unmapped_private_fields_never_leak(self):
        self.con.execute("UPDATE person SET raw_json=?, home_address=?, place_of_birth=? WHERE id=1",
                         (json.dumps({"custom_field": CANARY}), CANARY, CANARY))
        self.con.execute("UPDATE documents SET source_path=?, notes=?", (CANARY, CANARY))
        add_entry(self.con, title="Visible structured entry", include_long=0)
        add_entry(self.con, title="Hidden long-only entry", include_short=0)
        add_entry(self.con, section_key="unmapped_section", title=CANARY)
        text, result = self.render()
        self.assertIn("Visible structured entry", text)
        self.assertNotIn("Hidden long-only entry", text)
        self.assertNotIn(CANARY, text)
        self.assertEqual(result["manual_sections"], [])
        with zipfile.ZipFile(self.output) as archive:
            self.assertTrue(all(CANARY.encode() not in archive.read(name) for name in archive.namelist()))

    def test_dates_sort_newest_first_with_undated_last_and_no_raw_date_override(self):
        for start, end, title in (("", "", "Undated"), ("2018", "2020", "Earlier"),
                                  ("05/2025", "", "Latest"), ("01.04.2023", "", "Middle")):
            add_entry(self.con, start_date=start, end_date=end, title=title, raw_text="2099 | PRIVATE_DATE_CANARY")
        text, _ = self.render()
        self.assertLess(text.index("Latest"), text.index("Middle"))
        self.assertLess(text.index("Middle"), text.index("Earlier"))
        self.assertLess(text.index("Earlier"), text.index("Undated"))
        self.assertNotIn("2099", text)

    def test_pipe_description_retains_unique_details_without_repeating_fields(self):
        add_entry(self.con, title="Example doctorate", organization="Example University",
                  description="Example doctorate | Distinct thesis topic | Additional distinction",
                  start_date="2014", end_date="2014")
        text, _ = self.render()
        self.assertEqual(text.count("Example doctorate"), 1)
        self.assertIn("Distinct thesis topic", text)
        self.assertIn("Additional distinction", text)
        self.assertNotIn("2014–2014", text)

    def test_active_grants_respect_status_visibility_and_inclusive_dates(self):
        cases = (
            ("Current funded", "funded", "2025", "2027", 1, None, True),
            ("Ends today", "funded", "2025", "2026-09-07", 1, None, True),
            ("Year precision", "funded", "2025", "2026", 1, None, True),
            ("Month precision", "funded", "2025", "09/2026", 1, None, True),
            ("Ongoing funded", "funded", "2025", "Present", 1, None, True),
            ("Undated funded", "funded", "", "", 1, None, True),
            ("Legacy funded", None, "2025", "2027", 1, None, True),
            ("Legacy application", None, "2025", "2027", 1, "grant_application", False),
            ("Future grant", "funded", "2026-09-08", "2029", 1, None, False),
            ("Expired grant", "funded", "2025", "2026-09-06", 1, None, False),
            ("Expired month", "funded", "2025", "08/2026", 1, None, False),
            ("Expired year", "funded", "2020", "2025", 1, None, False),
            ("Explicitly past", "past", "2025", "2029", 1, None, False),
            ("Planned grant", "planned", "2025", "2029", 1, None, False),
            ("Submitted grant", "submitted", "2025", "2029", 1, None, False),
            ("Rejected grant", "rejected", "2025", "2029", 1, None, False),
            ("Hidden grant", "funded", "2025", "2029", 0, None, False),
            ("Invalid end month", "funded", "2025", "2026-13", 1, None, False),
        )
        for title, status, start, end, selected, subcategory, _ in cases:
            add_entry(self.con, section_key="funding", title=title, grant_status=status,
                      start_date=start, end_date=end, include_short=selected, subcategory=subcategory)
        text, result = self.render()
        for title, *_, included in cases:
            with self.subTest(grant=title):
                self.assertEqual(title in text, included)
        self.assertEqual(result["section_counts"]["funding"], 7)

    def test_five_selected_publications_keep_order_and_full_citations(self):
        for index in range(1, 8):
            add_publication(self.con, title=f"Paper number {index}", short_selected_order=8 - index,
                            raw_citation=f"Example J, Sample R, Model T, Demo A, Fiction B. Paper number {index}. "
                                         "Example Journal. 2026;12(3):45–56.",
                            doi=f"https://doi.org/10.0000/paper.{index}")
        add_publication(self.con, title="Suppressed selection", suppress_display=1, short_selected_order=0)
        add_publication(self.con, title="Unreviewed selection", category="preprint", short_selected_order=0)
        add_publication(self.con, title="Unselected publication", include_short=0)
        text, result = self.render()
        for index in range(3, 8):
            self.assertIn(f"Paper number {index}", text)
        for omitted in ("Paper number 1", "Paper number 2", "Suppressed selection",
                        "Unreviewed selection", "Unselected publication"):
            self.assertNotIn(omitted, text)
        self.assertLess(text.index("Paper number 7"), text.index("Paper number 6"))
        self.assertIn("Fiction B", text)
        self.assertIn("12(3):45–56", text)
        self.assertIn("https://doi.org/10.0000/paper.7", text)
        self.assertNotIn("https://doi.org/https://doi.org", text)
        self.assertEqual(result["section_counts"]["publications"], 5)

    def test_structured_citation_and_fallback_select_only_eligible_owner_papers(self):
        self.con.execute("UPDATE person SET full_name='Jane Example', display_name='Jane Example' WHERE id=1")
        for index in range(1, 7):
            add_publication(self.con, include_short=0, title=f"Owner paper {index}", year=str(2020 + index),
                            authors="Example J, Sample R, Model T, Demo A, Fiction B", doi=f"10.0000/owner.{index}")
        add_publication(self.con, include_short=0, title="Other authors", authors="Sample R, Model T")
        add_publication(self.con, include_short=0, title="Suppressed fallback", suppress_display=1)
        add_publication(self.con, include_short=0, title="Preprint fallback", category="preprint")
        text, result = self.render()
        for index in range(2, 7):
            self.assertIn(f"Owner paper {index}", text)
        for omitted in ("Owner paper 1", "Other authors", "Suppressed fallback", "Preprint fallback"):
            self.assertNotIn(omitted, text)
        self.assertIn("Fiction B", text)
        self.assertIn("Example Journal", text)
        self.assertEqual(result["section_counts"]["publications"], 5)
        self.con.execute("UPDATE publications SET include_short=1 WHERE title='Owner paper 2'")
        text, result = self.render()
        self.assertEqual(result["section_counts"]["publications"], 1)
        self.assertIn("■ Key papers", text)
        self.assertNotIn("■ Five key papers", text)
        self.assertNotIn("Owner paper 6", text)

    def test_empty_sections_and_unfilled_person_rows_disappear(self):
        text, result = self.render()
        self.assertEqual(text.strip(), "Curriculum vitae")
        self.assertEqual(result["rendered_items"], 0)
        self.assertEqual(result["rendered_sections"], [])
        self.assertNotIn("{{", text)
        self.con.execute("UPDATE person SET full_name='Jane Example' WHERE id=1")
        text, _ = self.render()
        self.assertIn("■ Personal data", text)
        self.assertIn("Name:", text)
        self.assertNotIn("ORCID:", text)
        self.assertNotIn("Professional address:", text)

    def test_page_numbers_are_dynamic_fields_and_update_in_every_render(self):
        populate_example(self.con)
        self.render()
        for path in (TEMPLATE_PATH, self.output):
            with self.subTest(path=path.name), zipfile.ZipFile(path) as archive:
                footer = etree.fromstring(archive.read("word/footer1.xml"))
                settings = etree.fromstring(archive.read("word/settings.xml"))
                fields = footer.findall(f".//{qn('w:fldSimple')}")
                self.assertEqual([field.get(qn("w:instr")) for field in fields], ["PAGE", "NUMPAGES"])
                self.assertTrue(all(field.get(qn("w:dirty")) == "true" for field in fields))
                self.assertIn("/", list(footer.itertext()))
                self.assertEqual(settings.find(qn("w:updateFields")).get(qn("w:val")), "true")
                document = Document(path)
                self.assertFalse(document.sections[0].different_first_page_header_footer)
                self.assertIsNone(document.sections[0]._sectPr.find(qn("w:pgNumType")))


class CompactResearchCvTemplatePrivacyTests(unittest.TestCase):
    def test_entire_template_archive_contains_only_generic_allowlisted_text(self):
        # An allowlist detects any source text without saving personal strings
        # from the reference CV in the repository or in test failure messages.
        members = {
            "[Content_Types].xml", "_rels/.rels", "docProps/app.xml", "docProps/core.xml",
            "word/document.xml", "word/_rels/document.xml.rels", "word/fontTable.xml",
            "word/settings.xml", "word/footer1.xml", "word/styles.xml", "word/numbering.xml",
            "word/webSettings.xml", "word/theme/theme1.xml",
        }
        body_text = {
            "Curriculum vitae", "■ Personal data", "Name:", "{{NAME}}", "Title:", "{{DEGREES}}",
            "ORCID:", "{{ORCID}}", "Current position:", "{{POSITION}}", "Professional address:",
            "{{ADDRESS}}", "■ Education/Degrees", "{{DATE}}", "{{EDUCATION}}",
            "■ Past and present positions", "{{POSITIONS}}", "■ Prizes and awards", "{{HONORS}}",
            "■ Activities in the Research System", "{{ACTIVITIES}}", "■ Active funding and grants",
            "{{FUNDING}}", "■ Five key papers", "{{PUBLICATIONS}}",
        }
        allowed_text = {
            "word/document.xml": body_text,
            "word/footer1.xml": {"1", "/", "Curriculum vitae"},
            "docProps/app.xml": {"VitaMine"},
            "docProps/core.xml": {"Compact Research CV", "VitaMine", "1", "2000-01-01T00:00:00Z"},
        }
        forbidden_elements = {
            "altChunk", "attachedTemplate", "bookmarkStart", "bookmarkEnd", "comment", "commentRangeStart",
            "commentRangeEnd", "commentReference", "customXml", "dataBinding", "del", "delText", "docVar",
            "docVars", "drawing", "hyperlink", "ins", "moveFrom", "moveTo", "object", "pict", "rsid",
            "rsids", "sdt", "smartTag", "trackRevisions",
        }
        relationship_types = {
            "extended-properties", "core-properties", "officeDocument", "fontTable", "settings",
            "footer", "styles", "numbering", "webSettings", "theme",
        }
        with zipfile.ZipFile(TEMPLATE_PATH) as archive:
            self.assertEqual(set(archive.namelist()), members)
            self.assertEqual(len(archive.namelist()), len(members))
            self.assertEqual(archive.comment, b"")
            for info in archive.infolist():
                with self.subTest(part=info.filename):
                    self.assertEqual(info.comment, b"")
                    root = etree.fromstring(archive.read(info.filename))
                    self.assertFalse(root.getroottree().docinfo.doctype)
                    self.assertFalse(root.xpath("//comment() | //processing-instruction()"))
                    actual = {value.strip() for value in root.itertext() if value.strip()}
                    self.assertTrue(actual <= allowed_text.get(info.filename, set()),
                                    "Unexpected text in sanitized template package")
                    if info.filename in allowed_text:
                        self.assertEqual(len(actual), len(allowed_text[info.filename]))
                    for element in root.iter():
                        self.assertNotIn(etree.QName(element).localname, forbidden_elements)
                        for attribute in element.attrib:
                            attribute_name = etree.QName(attribute).localname
                            self.assertFalse(attribute_name.lower().startswith("rsid"))
                            self.assertNotIn(attribute_name, {
                                "author", "initials", "descr", "description", "docId", "durableId",
                                "paraId", "textId", "email",
                            })
                    if info.filename.endswith(".rels"):
                        base = "word" if info.filename.startswith("word/") else ""
                        for relation in root:
                            self.assertNotEqual(relation.get("TargetMode"), "External")
                            self.assertIn(relation.get("Type").rsplit("/", 1)[-1], relationship_types)
                            self.assertIn(posixpath.normpath(posixpath.join(base, relation.get("Target"))), members)

    def test_metadata_is_fixed_and_contains_no_contact_identity_or_source_properties(self):
        with zipfile.ZipFile(TEMPLATE_PATH) as archive:
            core = etree.fromstring(archive.read("docProps/core.xml"))
            app = etree.fromstring(archive.read("docProps/app.xml"))
        expected = {
            "title": "Compact Research CV", "creator": "VitaMine", "lastModifiedBy": "VitaMine",
            "revision": "1", "created": "2000-01-01T00:00:00Z", "modified": "2000-01-01T00:00:00Z",
        }
        actual = {etree.QName(element).localname: element.text for element in core if element.text}
        self.assertEqual(actual, expected)
        self.assertEqual([(etree.QName(element).localname, element.text) for element in app],
                         [("Application", "VitaMine")])


if __name__ == "__main__":
    unittest.main()
