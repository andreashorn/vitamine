import os
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from vitamine.app import app
from vitamine.paths import create_blank_database


class DesignedCvExportTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        root = Path(self.directory.name)
        self.database = root / "designed-cv.vitamine"
        self.output = root / "output"
        create_blank_database(self.database)
        with sqlite3.connect(self.database) as con:
            con.execute(
                "UPDATE person SET full_name='Jane Example', display_name='Jane Example', position_title='Researcher', own_institution_name='Example University' WHERE id=1"
            )
            document_id = con.execute("SELECT id FROM documents LIMIT 1").fetchone()[0]
            con.execute(
                """
                INSERT INTO cv_entries (document_id, section_key, start_date, end_date, title, organization, description, raw_text)
                VALUES (?, 'academic_appointments', '2022', '', 'Group leader', 'Example University', 'Leads a research group', '2022 Group leader, Example University')
                """,
                (document_id,),
            )
            con.execute(
                """
                INSERT INTO publications (document_id, category, authors, title, venue, year, doi, raw_citation)
                VALUES (?, 'peer_reviewed', 'Example J', 'A current paper', 'Science', '2026', '10.1/example', 'Example J. A current paper. Science. 2026.')
                """,
                (document_id,),
            )
        self.environment = patch.dict(os.environ, {"VITAMINE_DB": str(self.database), "VITAMINE_OUTPUT": str(self.output)})
        self.environment.start()
        self.app_output = patch("vitamine.app.OUTPUT", self.output)
        self.paths_output = patch("vitamine.paths.OUTPUT", self.output)
        self.app_output.start()
        self.paths_output.start()
        self.preferences = {}
        self.read_preferences = patch("vitamine.app.read_preferences", side_effect=lambda: dict(self.preferences))
        self.write_preferences = patch("vitamine.app.write_preferences", side_effect=self._write_preferences)
        self.read_preferences.start()
        self.write_preferences.start()
        self.client = TestClient(app)

    def tearDown(self):
        self.client.close()
        self.write_preferences.stop()
        self.read_preferences.stop()
        self.paths_output.stop()
        self.app_output.stop()
        self.environment.stop()
        self.directory.cleanup()

    def _write_preferences(self, payload):
        self.preferences.clear()
        self.preferences.update(payload)

    def _export(self, format_id):
        with patch("vitamine.app.hosted_vitamine_plus_active", return_value=False):
            installed = self.client.post(f"/api/export-formats/{format_id}/install")
            self.assertEqual(installed.status_code, 200, installed.text)
            response = self.client.post(f"/api/actions/export/{format_id}?lang=en")
        self.assertEqual(response.status_code, 200, response.text)
        return self.output / Path(response.json()["docx_path"]).name

    def test_modern_publication_first_is_an_original_word_export(self):
        output = self._export("vitamine.modern-publication-first")
        self.assertTrue(output.exists())
        document = Document(output)
        text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        self.assertIn("Jane Example", text)
        self.assertIn("A current paper", text)
        self.assertIn("APPOINTMENTS", text)

    def test_r4ri_export_uses_the_four_editable_contributions(self):
        initial = self.client.get("/api/r4ri-contributions")
        self.assertEqual(initial.status_code, 200, initial.text)
        sections = initial.json()["sections"]
        self.assertEqual(len(sections), 4)
        sections[0]["body"] = "Created a reusable research method."
        saved = self.client.put("/api/r4ri-contributions", json={"sections": sections})
        self.assertEqual(saved.status_code, 200, saved.text)
        output = self._export("vitamine.r4ri-narrative")
        document = Document(output)
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Created a reusable research method.", text)
        self.assertIn("CONTRIBUTIONS TO THE GENERATION OF KNOWLEDGE", text)


if __name__ == "__main__":
    unittest.main()
