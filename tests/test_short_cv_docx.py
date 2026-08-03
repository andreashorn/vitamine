import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches

from vitamine.scripts.build_short_cv import (
    add_compact_heading,
    remove_table_borders,
    set_table_width,
)


class ShortCvDocxTests(unittest.TestCase):
    def test_table_geometry_uses_word_twips_and_no_visible_grid(self):
        document = Document()
        table = document.add_table(rows=1, cols=2)
        table.style = None
        widths = [Inches(1.45), Inches(5.55)]
        remove_table_borders(table)
        set_table_width(table, widths)

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "short.docx"
            document.save(path)
            rendered = Document(path)
            table = rendered.tables[0]
            expected = [str(int(width.twips)) for width in widths]
            grid = [column.get(qn("w:w")) for column in table._tbl.tblGrid]
            cells = [
                cell._tc.get_or_add_tcPr().find(qn("w:tcW")).get(qn("w:w"))
                for cell in table.rows[0].cells
            ]
            borders = table._tbl.tblPr.first_child_found_in("w:tblBorders")

            self.assertEqual(grid, expected)
            self.assertEqual(cells, expected)
            self.assertEqual(
                table._tbl.tblPr.first_child_found_in("w:tblW").get(qn("w:w")),
                str(sum(int(width.twips) for width in widths)),
            )
            self.assertTrue(all(edge.get(qn("w:val")) == "nil" for edge in borders))

    def test_section_heading_has_a_plain_top_rule(self):
        document = Document()
        add_compact_heading(document, "Selected Publications")
        paragraph = document.paragraphs[0]
        top = paragraph._p.pPr.find(qn("w:pBdr")).find(qn("w:top"))

        self.assertEqual(top.get(qn("w:val")), "single")
        self.assertEqual(top.get(qn("w:color")), "111111")


if __name__ == "__main__":
    unittest.main()
