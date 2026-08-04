#!/usr/bin/env python3
"""Build VitaMine's original designed Word CV exports."""

from __future__ import annotations

import argparse
import sqlite3
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from vitamine.paths import OUTPUT, active_db_path, output_ref
from vitamine.scripts.export_utils import sanitize_docx_compatibility_markup


R4RI_SECTIONS = (
    ("knowledge", "Contributions to the generation of knowledge", "Beiträge zur Gewinnung von Erkenntnissen"),
    ("people", "Contributions to the development of individuals and teams", "Beiträge zur Entwicklung von Personen und Teams"),
    ("research_community", "Contributions to the wider research and innovation community", "Beiträge zur Forschungs- und Innovationsgemeinschaft"),
    ("society", "Contributions to broader society and the economy", "Beiträge für Gesellschaft und Wirtschaft"),
)
MODERN_SECTIONS = (
    ("academic_appointments", "Appointments", "Positionen"),
    ("education", "Education", "Ausbildung"),
    ("funding", "Selected funding", "Ausgewählte Förderung"),
    ("honors", "Selected distinctions", "Ausgewählte Auszeichnungen"),
)


def clean(value: object | None) -> str:
    return str(value or "").strip()


def connect() -> sqlite3.Connection:
    con = sqlite3.connect(active_db_path())
    con.row_factory = sqlite3.Row
    return con


def ensure_r4ri_sections(con: sqlite3.Connection) -> None:
    con.execute(
        """
        CREATE TABLE IF NOT EXISTS r4ri_contribution_sections (
          section_key TEXT PRIMARY KEY,
          title TEXT NOT NULL,
          body TEXT NOT NULL DEFAULT '',
          title_de TEXT,
          body_de TEXT,
          updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        )
        """
    )
    for key, title, _de in R4RI_SECTIONS:
        con.execute("INSERT OR IGNORE INTO r4ri_contribution_sections (section_key, title) VALUES (?, ?)", (key, title))
    con.commit()


def set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    properties.append(shading)


def set_cell_border(cell, *, bottom: str | None = None) -> None:
    if not bottom:
        return
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    element = OxmlElement("w:bottom")
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), "6")
    element.set(qn("w:color"), bottom)
    borders.append(element)


def set_run(run, *, size: float = 10, bold: bool = False, color: str = "1C2733") -> None:
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def paragraph(container, value: str = "", *, size: float = 10, bold: bool = False, color: str = "1C2733", before: float = 0, after: float = 4):
    p = container.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.08
    if value:
        set_run(p.add_run(value), size=size, bold=bold, color=color)
    return p


def heading(container, value: str, *, color: str = "087E8B") -> None:
    p = paragraph(container, value.upper(), size=9.5, bold=True, color=color, before=10, after=4)
    p.paragraph_format.keep_with_next = True


def row_text(row: sqlite3.Row) -> str:
    period = "–".join(part for part in (clean(row["start_date"]), clean(row["end_date"])) if part)
    detail = " · ".join(part for part in (clean(row["title"]), clean(row["organization"]), clean(row["role"])) if part)
    description = clean(row["description"])
    return " — ".join(part for part in (period, detail, description) if part)


def publications(con: sqlite3.Connection, limit: int = 10) -> list[str]:
    rows = con.execute(
        """
        SELECT authors, title, venue, year, doi, raw_citation
        FROM publications
        WHERE COALESCE(suppress_display, 0)=0
        ORDER BY CAST(substr(COALESCE(year, '0'), 1, 4) AS INTEGER) DESC, id DESC
        LIMIT ?
        """,
        (limit,),
    ).fetchall()
    values = []
    for row in rows:
        citation = ". ".join(part.rstrip(".") for part in (clean(row["authors"]), clean(row["title"]), clean(row["venue"]), clean(row["year"])) if part)
        if row["doi"]:
            citation = f"{citation}. https://doi.org/{clean(row['doi'])}" if citation else f"https://doi.org/{clean(row['doi'])}"
        values.append(citation or clean(row["raw_citation"]))
    return [value for value in values if value]


def person(con: sqlite3.Connection) -> sqlite3.Row | None:
    return con.execute("SELECT * FROM person WHERE id=1").fetchone()


def document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.58)
    section.right_margin = Inches(0.58)
    return doc


def add_modern_cv(con: sqlite3.Connection, lang: str) -> Document:
    doc = document()
    individual = person(con)
    name = clean(individual["display_name"] if individual and "display_name" in individual.keys() else "") or clean(individual["full_name"] if individual else "") or "Curriculum Vitae"
    title = clean(individual["position_title"] if individual and "position_title" in individual.keys() else "")
    institution = clean(individual["own_institution_name"] if individual and "own_institution_name" in individual.keys() else "")
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    rail, body = table.rows[0].cells
    rail.width, body.width = Inches(1.85), Inches(5.65)
    rail.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    body.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_shading(rail, "123047")
    rail._tc.tcPr.tcW.set(qn("w:w"), "2664")
    body._tc.tcPr.tcW.set(qn("w:w"), "8136")
    paragraph(rail, name, size=20, bold=True, color="FFFFFF", after=8)
    if title:
        paragraph(rail, title, size=9.5, color="DCE8ED", after=10)
    if institution:
        paragraph(rail, institution, size=9.5, color="DCE8ED", after=14)
    heading(rail, "Profile" if lang == "en" else "Profil", color="74C7C7")
    narrative = con.execute("SELECT body, body_de FROM narrative_reports WHERE id=1").fetchone()
    profile = clean(narrative["body_de"] if narrative and lang == "de" and clean(narrative["body_de"]) else narrative["body"] if narrative else "")
    paragraph(rail, profile or ("Add a research profile in the Narrative tab." if lang == "en" else "Fügen Sie im Bereich Narrative ein Forschungsprofil hinzu."), size=8.5, color="DCE8ED", after=12)
    heading(rail, "Identifiers" if lang == "en" else "Kennungen", color="74C7C7")
    for field in ("orcid_id", "email", "website"):
        value = clean(individual[field] if individual and field in individual.keys() else "")
        if value:
            paragraph(rail, value, size=8.5, color="DCE8ED", after=4)
    paragraph(body, "CURRICULUM VITAE" if lang == "en" else "LEBENSLAUF", size=9.5, bold=True, color="087E8B", after=3)
    paragraph(body, name, size=26, bold=True, color="123047", after=3)
    if title or institution:
        paragraph(body, " · ".join(part for part in (title, institution) if part), size=11, color="51616D", after=8)
    for key, english, german in MODERN_SECTIONS:
        rows = con.execute("SELECT * FROM cv_entries WHERE section_key=? ORDER BY start_date DESC, id DESC LIMIT 8", (key,)).fetchall()
        if not rows:
            continue
        heading(body, english if lang == "en" else german)
        for row in rows:
            paragraph(body, row_text(row), size=9.5, after=3)
    selected = publications(con, 12)
    if selected:
        heading(body, "Selected publications" if lang == "en" else "Ausgewählte Publikationen")
        for index, value in enumerate(selected, 1):
            paragraph(body, f"{index}. {value}", size=8.8, after=3)
    return doc


def add_r4ri_cv(con: sqlite3.Connection, lang: str) -> Document:
    ensure_r4ri_sections(con)
    doc = document()
    individual = person(con)
    name = clean(individual["display_name"] if individual and "display_name" in individual.keys() else "") or clean(individual["full_name"] if individual else "") or "Curriculum Vitae"
    title = clean(individual["position_title"] if individual and "position_title" in individual.keys() else "")
    p = paragraph(doc, "RÉSUMÉ FOR RESEARCH AND INNOVATION" if lang == "en" else "LEBENSLAUF FÜR FORSCHUNG UND INNOVATION", size=10, bold=True, color="087E8B", after=4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = paragraph(doc, name, size=23, bold=True, color="123047", after=3)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if title:
        p = paragraph(doc, title, size=10.5, color="51616D", after=10)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice = "VitaMine drafting format — check the current funder instructions before submission." if lang == "en" else "VitaMine-Entwurf — prüfen Sie vor Einreichung die aktuellen Vorgaben des Förderers."
    p = paragraph(doc, notice, size=8.5, color="51616D", after=10)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    entries = {row["section_key"]: row for row in con.execute("SELECT * FROM r4ri_contribution_sections").fetchall()}
    for key, english, german in R4RI_SECTIONS:
        row = entries.get(key)
        body = clean(row["body_de"] if row and lang == "de" and clean(row["body_de"]) else row["body"] if row else "")
        title_value = clean(row["title_de"] if row and lang == "de" and clean(row["title_de"]) else row["title"] if row else "") or (english if lang == "en" else german)
        heading(doc, title_value)
        paragraph(doc, body or ("Add this contribution in the Narrative tab." if lang == "en" else "Fügen Sie diesen Beitrag im Bereich Narrative hinzu."), size=10, after=8)
    selected = publications(con, 10)
    if selected:
        heading(doc, "Selected outputs" if lang == "en" else "Ausgewählte Ergebnisse")
        for index, value in enumerate(selected, 1):
            paragraph(doc, f"{index}. {value}", size=9, after=3)
    return doc


def build(format_name: str, lang: str) -> Path:
    con = connect()
    try:
        doc = add_modern_cv(con, lang) if format_name == "modern" else add_r4ri_cv(con, lang)
    finally:
        con.close()
    OUTPUT.mkdir(parents=True, exist_ok=True)
    suffix = "_de" if lang == "de" else ""
    filename = "modern_publication_first" if format_name == "modern" else "narrative_research_cv"
    path = OUTPUT / f"{filename}{suffix}.docx"
    doc.core_properties.title = "Modern Publication-First CV" if format_name == "modern" else "Narrative Research CV"
    doc.core_properties.author = "VitaMine"
    doc.save(path)
    sanitize_docx_compatibility_markup(path)
    return path


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--format", choices=("modern", "r4ri"), required=True)
    parser.add_argument("--lang", choices=("en", "de"), default="en")
    args = parser.parse_args()
    print(f"docx: output/{output_ref(build(args.format, args.lang))}")
