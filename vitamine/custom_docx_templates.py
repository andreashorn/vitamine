"""Import Word CVs as private, reusable VitaMine export templates.

The uploaded document remains a Word-native package.  We replace recognized
person and section content with tokens while retaining its page geometry,
styles, tables, headers, footers, and images.  Export then pours the current
content from one of VitaMine's deterministic builders back into those slots.
"""

from __future__ import annotations

import copy
import hashlib
import io
import json
import re
import sqlite3
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterable

from docx import Document
from docx.document import Document as DocumentObject
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.oxml.table import CT_Row, CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Row
from docx.text.paragraph import Paragraph
from lxml import etree

from .llm_routing import settings_for_llm_task


MAX_TEMPLATE_BYTES = 20 * 1024 * 1024
MAX_TEMPLATE_UNCOMPRESSED_BYTES = 120 * 1024 * 1024
MAX_TEMPLATE_PARTS = 2_000

CONTENT_PROFILES = {"long", "short", "one_page", "biosketch"}
PERSON_FIELDS = (
    "full_name",
    "display_name",
    "degrees",
    "position_title",
    "office_address",
    "home_address",
    "work_phone",
    "work_email",
    "place_of_birth",
    "era_commons",
    "orcid_id",
    "own_institution_name",
)

SECTION_ALIASES: dict[str, tuple[str, ...]] = {
    "qualifications_and_career": ("qualifications and career",),
    "research_system_activities": ("activities in the research system",),
    "research_experience": (
        "research experience", "research employment", "research positions",
        "scientific experience", "wissenschaftliche erfahrung",
    ),
    "research_skills": (
        "relevant research skills", "research skills", "scientific skills",
        "technical skills", "laboratory skills", "methodological skills",
        "forschungsmethoden", "wissenschaftliche kompetenzen",
    ),
    "education": (
        "education", "education and training", "education/training", "academic education",
        "ausbildung", "studium und ausbildung", "akademische ausbildung", "qualifications",
    ),
    "postdoctoral_training": (
        "postdoctoral training", "postdoctoral education", "postdoctoral fellowships",
        "postdoktorale ausbildung", "postdoc training",
    ),
    "academic_appointments": (
        "academic appointments", "faculty academic appointments", "scientific appointments",
        "positions and scientific appointments", "positions and appointments", "academic positions",
        "akademische berufungen", "akademische positionen", "positionen und berufungen",
    ),
    "hospital_appointments": (
        "hospital appointments", "appointments at hospitals/affiliated institutions",
        "clinical appointments", "klinische positionen",
    ),
    "professional_positions": (
        "professional positions", "other professional positions", "employment", "experience",
        "work experience", "career history", "additional relevant experience",
        "beruflicher werdegang", "berufserfahrung",
    ),
    "committee_service": (
        "committee service", "committee membership", "committees", "gremienarbeit",
    ),
    "professional_societies": (
        "professional societies", "professional memberships", "memberships", "mitgliedschaften",
        "fachgesellschaften",
    ),
    "grant_review": (
        "grant review activities", "grant review", "review panels", "gutachtertätigkeiten",
    ),
    "editorial_activities": (
        "editorial activities", "editorial service", "journal service", "editoriale tätigkeiten",
    ),
    "honors": (
        "honors", "honours", "honors and prizes", "honors and awards", "awards",
        "awards and honors", "selected honors", "distinctions", "academic distinctions", "auszeichnungen",
        "ausgewählte auszeichnungen", "auszeichnungen und preise",
    ),
    "funding": (
        "research funding", "funding", "selected funding", "grants", "grant support", "funded projects",
        "report of funded and unfunded projects", "forschungsförderung", "drittmittel",
        "ausgewählte forschungsförderung",
    ),
    "teaching": (
        "teaching", "teaching experience", "teaching activities", "teaching of students in courses", "lehre",
    ),
    "mentoring": (
        "mentoring", "supervision", "research supervisory and training responsibilities",
        "supervision of researchers in early career phases selection",
        "trainees", "nachwuchsförderung", "betreuung und ausbildung",
    ),
    "invited_presentations": (
        "invited presentations", "invited lectures", "invited talks", "presentations",
        "selected presentations", "eingeladene vorträge", "ausgewählte vorträge", "vorträge",
    ),
    "clinical_activities": (
        "clinical activities", "clinical activities and innovations", "klinische tätigkeiten",
    ),
    "education_innovations": (
        "teaching and education innovations", "education innovations", "lehrinnovationen",
    ),
    "community_service": (
        "community service", "service to the community", "gesellschaftliches engagement",
    ),
    "publications": (
        "publications", "selected publications", "peer-reviewed publications",
        "peer-reviewed scholarship in print or other media", "report of scholarship",
        "bibliography", "scientific output", "works", "publikationen", "ausgewählte publikationen",
        "veröffentlichungen",
    ),
    "dfg_category_a": (
        "category a articles in peer-reviewed journals contributions to peer-reviewed conferences or to anthology volumes and book publications",
    ),
    "dfg_category_b": ("category b any other form of published results",),
    "supplementary_career_information": ("supplementary career information",),
    "dfg_data_protection": ("data protection and consent to the processing of optional data",),
    "dfg_scientific_results": ("scientific results",),
    "dfg_other_information": ("other information",),
    "personal_statement": (
        "personal statement", "summary statement", "research profile", "profile", "profil",
    ),
    "contributions": (
        "contributions to science", "contributions", "contributions to research",
        "wissenschaftliche beiträge",
    ),
    "narrative_report": (
        "narrative report", "research narrative", "narrativer bericht",
    ),
    "conference_papers": (
        "conference papers", "conference contributions", "conference abstracts",
        "kongressbeiträge", "konferenzbeiträge",
    ),
}

SOURCE_SECTION_FALLBACKS: dict[str, tuple[str, ...]] = {
    "qualifications_and_career": (
        "education", "postdoctoral_training", "academic_appointments",
        "hospital_appointments", "professional_positions",
    ),
    "research_system_activities": (
        "editorial_activities", "grant_review", "committee_service",
        "professional_societies", "community_service", "invited_presentations",
    ),
    "dfg_category_a": ("publications",),
    "research_experience": (
        "postdoctoral_training", "academic_appointments", "hospital_appointments",
        "professional_positions", "clinical_activities",
    ),
    "research_skills": ("contributions", "personal_statement", "narrative_report"),
    # Conference papers are deliberately not a synonym for invited talks. If
    # VitaMine has no matching data, the source block becomes a manual section.
    "conference_papers": (),
}

MANUAL_SECTION_PLACEHOLDER = "[Please fill this section manually.]"
MANUAL_ONLY_SECTION_KEYS = {
    "conference_papers", "dfg_category_b", "supplementary_career_information",
    "dfg_other_information",
}
PRESERVED_SECTION_KEYS = {"dfg_data_protection", "dfg_scientific_results"}

STATIC_LABELS = {
    "dates", "date", "years", "year", "degree", "field of study", "institution",
    "institution and location", "position", "title", "organization", "details", "role",
    "completion date", "current", "past", "local", "national", "international",
    "ad hoc reviewer", "other editorial roles", "activities", "research investigations",
    "other scholarship", "patents", "books / book chapters", "books / chapters", "preprints",
    "manuscripts in preparation", "poster presentations", "achievements",
    "other formally supervised trainees", "selected projects relevant to this application",
    "selected citations that highlight experience and qualifications",
}

LLM_ANALYSIS_SCHEMA: dict[str, Any] = {
    "type": "object",
    "additionalProperties": False,
    "required": ["content_profile", "confidence", "heading_mappings"],
    "properties": {
        "content_profile": {"type": "string", "enum": ["long", "short", "one_page", "biosketch"]},
        "confidence": {"type": "string", "enum": ["high", "medium", "low"]},
        "heading_mappings": {
            "type": "array",
            "maxItems": 80,
            "items": {
                "type": "object",
                "additionalProperties": False,
                "required": ["source_text", "section_key"],
                "properties": {
                    "source_text": {"type": "string"},
                    "section_key": {"type": "string", "enum": sorted(SECTION_ALIASES)},
                },
            },
        },
    },
}


@dataclass
class Unit:
    index: int
    kind: str
    element: Any
    parent: Any
    paragraphs: list[Paragraph]
    cell_paragraphs: list[list[Paragraph]]

    @property
    def values(self) -> list[str]:
        if self.kind == "paragraph":
            return [clean_text(self.paragraphs[0].text)]
        return [clean_text("\n".join(paragraph.text for paragraph in paragraphs)) for paragraphs in self.cell_paragraphs]

    @property
    def text(self) -> str:
        return clean_text(" | ".join(value for value in self.values if value))


def clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def normalized_heading(value: str) -> str:
    text = clean_text(value).casefold()
    text = re.sub(r"^[a-z0-9ivx]+[.)]\s+", "", text)
    text = text.rstrip(":.- ")
    text = re.sub(r"[^\wäöüß/&+ -]+", "", text)
    return re.sub(r"\s+", " ", text).strip()


def unique_cells(row: _Row) -> list[Any]:
    cells: list[Any] = []
    seen: set[int] = set()
    for cell in row.cells:
        marker = id(cell._tc)
        if marker not in seen:
            seen.add(marker)
            cells.append(cell)
    return cells


def document_units(document: DocumentObject) -> list[Unit]:
    units: list[Unit] = []
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            paragraph = Paragraph(child, document)
            units.append(Unit(len(units), "paragraph", child, document, [paragraph], []))
        elif isinstance(child, CT_Tbl):
            table = Table(child, document)
            for row in table.rows:
                cells = unique_cells(row)
                cell_paragraphs = [list(cell.paragraphs) for cell in cells]
                paragraphs = [paragraph for group in cell_paragraphs for paragraph in group]
                units.append(Unit(len(units), "row", row._tr, table, paragraphs, cell_paragraphs))
    return units


def paragraphs_in_container(container: Any) -> Iterable[Paragraph]:
    for paragraph in getattr(container, "paragraphs", []):
        yield paragraph
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in unique_cells(row):
                yield from paragraphs_in_container(cell)


def all_document_paragraphs(document: DocumentObject) -> Iterable[Paragraph]:
    yield from paragraphs_in_container(document)
    # Accessing ``section.header.part`` creates empty header/footer package
    # parts in python-docx.  Traverse only relationships that already existed
    # in the uploaded document so a plain CV remains a plain CV.
    seen_elements: set[int] = set()
    for relationship in document.part.rels.values():
        if relationship.reltype not in {RT.HEADER, RT.FOOTER}:
            continue
        root = relationship.target_part.element
        for element in root.iter(qn("w:p")):
            marker = id(element)
            if marker in seen_elements:
                continue
            seen_elements.add(marker)
            yield Paragraph(element, document)


def paragraph_is_bold(paragraph: Paragraph) -> bool:
    runs = [run for run in paragraph.runs if clean_text(run.text)]
    return bool(runs) and all(bool(run.bold) for run in runs)


def probable_heading(unit: Unit) -> bool:
    text = unit.text
    if not text or len(text) > 180:
        return False
    normalized = normalized_heading(text)
    if normalized in STATIC_LABELS:
        return True
    if any((paragraph.style and str(paragraph.style.name).casefold().startswith("heading")) for paragraph in unit.paragraphs):
        return True
    if unit.paragraphs and all(paragraph_is_bold(paragraph) for paragraph in unit.paragraphs if clean_text(paragraph.text)):
        return True
    letters = [character for character in text if character.isalpha()]
    return bool(letters) and len(text) <= 100 and sum(character.isupper() for character in letters) / len(letters) > 0.82


def paragraph_heading_level(paragraph: Paragraph) -> int | None:
    name = str(paragraph.style.name if paragraph.style else "")
    match = re.match(r"heading\s+(\d+)", name, flags=re.IGNORECASE)
    return int(match.group(1)) if match else None


def unit_heading_level(unit: Unit) -> int | None:
    levels = [level for paragraph in unit.paragraphs if (level := paragraph_heading_level(paragraph)) is not None]
    return min(levels) if levels else None


def mostly_uppercase(value: str) -> bool:
    letters = [character for character in value if character.isalpha()]
    return bool(letters) and sum(character.isupper() for character in letters) / len(letters) > 0.82


def probable_section_heading(unit: Unit) -> bool:
    """Return whether *unit* is a section boundary rather than an entry title.

    CVs commonly use Heading 2 for individual jobs.  Treating every styled
    heading as a section caused those job titles, their bullets, and later
    sections to bleed into one another.  Heading 1, explicit aliases, and
    short all-caps labels are much safer section signals.
    """

    if alias_section_key(unit.text):
        return True
    if unit_heading_level(unit) == 1:
        return True
    words = re.findall(r"[A-Za-zÄÖÜäöüß]+", unit.text)
    return bool(unit.text) and len(unit.text) <= 100 and len(words) >= 2 and mostly_uppercase(unit.text)


def alias_section_key(value: str) -> str | None:
    normalized = normalized_heading(value)
    if not normalized:
        return None
    # Exact matches must win before fuzzy suffix matches.  Otherwise
    # "Research Experience" is swallowed by the generic alias "Experience".
    for section_key, aliases in SECTION_ALIASES.items():
        if normalized in aliases:
            return section_key
    for section_key, aliases in SECTION_ALIASES.items():
        for alias in aliases:
            if len(alias) >= 12 and " " in alias and (
                normalized.startswith(f"{alias} ") or normalized.endswith(f" {alias}")
            ):
                return section_key
    return None


def mapped_section_key(value: str, mappings: dict[str, str]) -> str | None:
    normalized = normalized_heading(value)
    return mappings.get(normalized) or alias_section_key(value)


def plausible_llm_heading_mapping(source_text: str, section_key: str) -> bool:
    """Reject broad semantic guesses that would invent a database meaning."""

    normalized = normalized_heading(source_text)
    if section_key != "research_skills":
        return True
    research_markers = (
        "research", "scientific", "laboratory", "labor", "technical", "method",
        "experimental", "computational", "bioinformatic", "statistical",
        "wissenschaft", "forsch", "methoden",
    )
    return any(marker in normalized for marker in research_markers)


def page_count_from_docx(data: bytes) -> int | None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            raw = archive.read("docProps/app.xml")
        root = etree.fromstring(raw)
        pages = root.xpath("//*[local-name()='Pages']/text()")
        return int(pages[0]) if pages and int(pages[0]) > 0 else None
    except (KeyError, ValueError, zipfile.BadZipFile, etree.XMLSyntaxError):
        return None


def validate_docx_bytes(data: bytes) -> None:
    if not data:
        raise ValueError("The Word template is empty.")
    if len(data) > MAX_TEMPLATE_BYTES:
        raise ValueError("The Word template is larger than 20 MB.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_TEMPLATE_PARTS:
                raise ValueError("The Word template contains too many package parts.")
            if sum(info.file_size for info in infos) > MAX_TEMPLATE_UNCOMPRESSED_BYTES:
                raise ValueError("The expanded Word template is too large.")
            if any(info.flag_bits & 0x1 for info in infos):
                raise ValueError("Password-protected Word templates are not supported.")
            names = {info.filename for info in infos}
            if "word/document.xml" not in names or "[Content_Types].xml" not in names:
                raise ValueError("This file is not a readable DOCX document.")
            if any("vbaproject" in name.casefold() for name in names):
                raise ValueError("Macro-enabled Word documents are not supported.")
    except zipfile.BadZipFile as exc:
        raise ValueError("This file is not a readable DOCX document.") from exc


def deterministic_profile(text: str, *, page_count: int | None, heading_count: int) -> tuple[str, str]:
    normalized = normalized_heading(text)
    words = len(re.findall(r"\b\w+\b", text))
    biosketch_markers = sum(
        marker in normalized
        for marker in ("biographical sketch", "biosketch", "personal statement", "contributions to science", "era commons")
    )
    if biosketch_markers >= 2 or "contributions to science" in normalized:
        return "biosketch", "biosketch structure"
    if (page_count == 1 and words <= 1_600 and heading_count <= 10) or (words <= 750 and heading_count <= 8):
        return "one_page", "one-page or very compact structure"
    # A curated 2-5 page CV can have many visible headings because role titles
    # often use Heading 2.  Length, not raw heading count, distinguishes it
    # from VitaMine's comprehensive long profile.
    if (page_count is not None and page_count >= 6) or words >= 2_400:
        return "long", "comprehensive multi-section structure"
    return "short", "curated multi-page academic CV structure"


def llm_prompt(units: list[Unit], page_count: int | None, deterministic: str) -> str:
    candidates = []
    for unit in units:
        text = unit.text
        if text and len(text) <= 220 and probable_section_heading(unit):
            candidates.append(text)
        if len(candidates) >= 120:
            break
    excerpt = "\n".join(unit.text for unit in units if unit.text)[:14_000]
    allowed = ", ".join(sorted(SECTION_ALIASES))
    return f"""Analyze the structure of this academic CV Word document for a reusable export template.

Classify its content shape as exactly one of long, short, one_page, or biosketch.
Use short for a curated 2-5 page CV with selected entries. Use long only for a
comprehensive CV that is generally at least 6 pages or roughly 2,400 words.
Map only headings that appear verbatim in the candidate list to the closest allowed section key.
Leave a heading unmapped when none of the allowed section keys describes it; VitaMine will preserve
that source section with a manual-fill placeholder instead of inventing content for it.
Generic skills, courses, software, languages, interests, references, and referees are not
research_skills unless the heading itself clearly identifies research, scientific, laboratory,
technical, methodological, experimental, computational, bioinformatics, or statistical skills.
Do not treat a person's name, institution, degree, date, or CV title as a section heading.
Allowed section keys: {allowed}

Document metadata:
- recorded pages: {page_count if page_count is not None else 'unknown'}
- deterministic first pass: {deterministic}
- body units: {len(units)}

Candidate headings:
{json.dumps(candidates, ensure_ascii=False)}

Document excerpt:
{excerpt}
"""


def section_slots(units: list[Unit], mappings: dict[str, str]) -> tuple[list[dict[str, Any]], list[str]]:
    slots: list[dict[str, Any]] = []
    mapped_sections: list[str] = []
    current: str | None = None
    for unit in units:
        section_key = mapped_section_key(unit.text, mappings)
        if section_key:
            current = section_key
            if section_key not in mapped_sections:
                mapped_sections.append(section_key)
            continue
        if not current or not unit.text or not content_unit(unit):
            continue
        slots.append(
            {
                "unit_index": unit.index,
                "section_key": current,
                "kind": unit.kind,
                "columns": max(1, len(unit.values)),
            }
        )
    return slots, mapped_sections


def masthead_identity_slots(units: list[Unit], mappings: dict[str, str]) -> list[dict[str, Any]]:
    """Identify the visible name/contact block before the first CV section."""

    first_section = next(
        (unit.index for unit in units if mapped_section_key(unit.text, mappings)),
        len(units),
    )
    candidates = [
        unit for unit in units[:first_section]
        if unit.text and normalized_heading(unit.text) not in {"cv", "curriculum vitae"}
    ]
    if not candidates:
        return []
    dfg_roles = {
        "title": "degrees",
        "first name": "first_name",
        "name": "last_name",
        "last name": "last_name",
        "current position": "position_title",
        "current institutions/sites country": "institution",
        "identifiers/orcid": "orcid_id",
    }
    dfg_slots: list[dict[str, Any]] = []
    for unit in candidates:
        if unit.kind != "row" or len(unit.values) < 2:
            continue
        role = dfg_roles.get(normalized_heading(unit.values[0]))
        if role:
            dfg_slots.append({"unit_index": unit.index, "role": role, "cell_index": 1})
    if len(dfg_slots) >= 3:
        return dfg_slots
    slots: list[dict[str, Any]] = []
    name_assigned = False
    for unit in candidates:
        text = unit.text
        if not name_assigned and len(text) <= 100 and not re.search(r"@|https?://|\d{3,}", text):
            role = "display_name"
            name_assigned = True
        elif re.search(r"@|\b(?:tel|phone|fax)\b|\d{4,}|\b(?:road|street|strasse|straße|avenue|university|department)\b", text, re.IGNORECASE):
            role = "contact_line"
        else:
            role = "position_title"
        slots.append({"unit_index": unit.index, "role": role})
    return slots


def entry_anchor(unit: Unit) -> bool:
    """Return whether a source paragraph is the reusable anchor for one record."""

    if unit.kind != "paragraph" or not unit.text:
        return False
    paragraph = unit.paragraphs[0]
    numbered = paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None
    level = unit_heading_level(unit)
    if level is not None and level >= 2 and not mostly_uppercase(unit.text):
        return True
    return not numbered and len(unit.text) <= 220 and paragraph_is_bold(paragraph)


def semantic_section_blueprint(
    units: list[Unit], mappings: dict[str, str]
) -> tuple[list[dict[str, Any]], list[str]]:
    """Describe source sections without allowing content to bleed across headings."""

    first_section = next(
        (unit.index for unit in units if mapped_section_key(unit.text, mappings)),
        None,
    )
    if first_section is None:
        return [], []

    sections: list[dict[str, Any]] = []
    mapped_sections: list[str] = []
    current: dict[str, Any] | None = None
    unknown_counter = 0
    for unit in units[first_section:]:
        section_key = mapped_section_key(unit.text, mappings)
        if section_key:
            current = {
                "section_key": section_key,
                "source_heading": unit.text,
                "heading_unit_index": unit.index,
                "body_unit_indices": [],
            }
            sections.append(current)
            if section_key not in mapped_sections:
                mapped_sections.append(section_key)
            continue
        if probable_section_heading(unit):
            unknown_counter += 1
            current = {
                "section_key": f"__unmapped_{unknown_counter}",
                "source_heading": unit.text,
                "heading_unit_index": unit.index,
                "body_unit_indices": [],
            }
            sections.append(current)
            continue
        if current is not None:
            current["body_unit_indices"].append(unit.index)

    by_index = {unit.index: unit for unit in units}
    for section in sections:
        body_units = [by_index[index] for index in section["body_unit_indices"] if index in by_index]
        content_units = [unit for unit in body_units if unit.text and content_unit(unit)]
        anchors = [unit for unit in content_units if entry_anchor(unit)]
        numbered_units = [
            unit for unit in content_units
            if unit.kind == "paragraph"
            and unit.paragraphs[0]._p.pPr is not None
            and unit.paragraphs[0]._p.pPr.numPr is not None
        ]
        record_units = anchors or numbered_units or content_units
        record_indices = {unit.index for unit in record_units}
        section["record_unit_indices"] = [unit.index for unit in record_units]
        section["discard_unit_indices"] = [
            unit.index for unit in body_units if unit.text and unit.index not in record_indices
        ]
    return sections, mapped_sections


def content_unit(unit: Unit) -> bool:
    text = unit.text
    normalized = normalized_heading(text)
    if not text or normalized in STATIC_LABELS:
        return False
    normalized_values = [normalized_heading(value) for value in unit.values if clean_text(value)]
    if normalized_values and all(value in STATIC_LABELS for value in normalized_values):
        return False
    if mapped_section_key(text, {}):
        return False
    return True


def set_paragraph_text(paragraph: Paragraph, value: str) -> None:
    run_properties = None
    for run in paragraph.runs:
        if run._r.rPr is not None:
            run_properties = copy.deepcopy(run._r.rPr)
            break
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)
    run = paragraph.add_run(value)
    if run_properties is not None:
        run._r.insert(0, run_properties)


def replace_paragraph_text(paragraph: Paragraph, old: str, new: str) -> bool:
    text = paragraph.text
    if not old or old.casefold() not in text.casefold():
        return False
    match = re.search(re.escape(old), text, flags=re.IGNORECASE)
    if not match:
        return False
    set_paragraph_text(paragraph, text[:match.start()] + new + text[match.end():])
    return True


def set_unit_values(unit: Unit, values: list[str]) -> None:
    values = [clean_text(value) for value in values]
    if unit.kind == "paragraph":
        populated = [value for value in values if value]
        if len(populated) == 2 and re.fullmatch(r"\d+[.)]", populated[0]):
            rendered = populated[1]
        elif len(populated) >= 2 and re.fullmatch(
            r"(?:\d{1,2}[./-]){0,2}\d{2,4}(?:\s*[-–—]\s*(?:present|current|heute|\d{2,4})?)?",
            populated[0],
            flags=re.IGNORECASE,
        ):
            rendered = f"{' · '.join(populated[1:])} ({populated[0]})"
        else:
            rendered = " · ".join(populated)
        set_paragraph_text(unit.paragraphs[0], rendered)
        return
    count = len(unit.cell_paragraphs)
    if count <= 1:
        adapted = [" · ".join(value for value in values if value)]
    elif len(values) <= count:
        adapted = values + [""] * (count - len(values))
    else:
        adapted = values[: count - 1] + [" · ".join(value for value in values[count - 1:] if value)]
    for paragraphs, value in zip(unit.cell_paragraphs, adapted):
        if paragraphs:
            set_paragraph_text(paragraphs[0], value)
            for paragraph in paragraphs[1:]:
                set_paragraph_text(paragraph, "")


def set_identity_slot_value(unit: Unit, slot: dict[str, Any], value: str) -> None:
    cell_index = slot.get("cell_index")
    if unit.kind == "row" and isinstance(cell_index, int) and 0 <= cell_index < len(unit.cell_paragraphs):
        paragraphs = unit.cell_paragraphs[cell_index]
        if paragraphs:
            set_paragraph_text(paragraphs[0], value)
            for paragraph in paragraphs[1:]:
                set_paragraph_text(paragraph, "")
        return
    set_unit_values(unit, [value])


def person_values(con: sqlite3.Connection) -> dict[str, str]:
    row = con.execute("SELECT * FROM person WHERE id=1").fetchone()
    if row is None:
        return {}
    keys = set(row.keys())
    return {field: clean_text(row[field]) for field in PERSON_FIELDS if field in keys and clean_text(row[field])}


def skeletonize_person(document: DocumentObject, units: list[Unit], values: dict[str, str], mappings: dict[str, str]) -> list[str]:
    first_heading = next((unit.index for unit in units if mapped_section_key(unit.text, mappings)), len(units))
    body_paragraph_ids = {
        id(paragraph._p)
        for unit in units[:first_heading]
        for paragraph in unit.paragraphs
    }
    existing_part_paragraph_ids = {
        id(paragraph._p)
        for paragraph in all_document_paragraphs(document)
        if id(paragraph._p) not in {
            id(body_paragraph._p)
            for body_unit in units
            for body_paragraph in body_unit.paragraphs
        }
    }
    eligible_paragraph_ids = body_paragraph_ids | existing_part_paragraph_ids
    fields: list[str] = []
    candidates = sorted(values.items(), key=lambda item: len(item[1]), reverse=True)
    for paragraph in all_document_paragraphs(document):
        if id(paragraph._p) not in eligible_paragraph_ids:
            continue
        for field, value in candidates:
            if len(value) < 4 or field in fields and f"{{{{VITAMINE_PERSON_{field.upper()}}}}}" in paragraph.text:
                continue
            token = f"{{{{VITAMINE_PERSON_{field.upper()}}}}}"
            if replace_paragraph_text(paragraph, value, token) and field not in fields:
                fields.append(field)
    return fields


def skeletonize_slots(units: list[Unit], slots: list[dict[str, Any]]) -> None:
    by_index = {unit.index: unit for unit in units}
    counters: dict[str, int] = defaultdict(int)
    for slot in slots:
        unit = by_index.get(int(slot["unit_index"]))
        if unit is None:
            continue
        section_key = str(slot["section_key"])
        counters[section_key] += 1
        token = f"{{{{VITAMINE_{section_key.upper()}_{counters[section_key]}}}}}"
        set_unit_values(unit, [token] * max(1, int(slot.get("columns") or 1)))


def skeletonize_semantic_blueprint(
    units: list[Unit], identity_slots: list[dict[str, Any]], sections: list[dict[str, Any]]
) -> None:
    by_index = {unit.index: unit for unit in units}
    for slot in identity_slots:
        unit = by_index.get(int(slot["unit_index"]))
        if unit is not None:
            set_identity_slot_value(unit, slot, f"{{{{VITAMINE_IDENTITY_{str(slot['role']).upper()}}}}}")
    for section_number, section in enumerate(sections, 1):
        if str(section.get("section_key") or "") in PRESERVED_SECTION_KEYS:
            continue
        indices = list(section.get("record_unit_indices") or []) + list(section.get("discard_unit_indices") or [])
        for item_number, unit_index in enumerate(indices, 1):
            unit = by_index.get(int(unit_index))
            if unit is not None:
                set_unit_values(unit, [f"{{{{VITAMINE_SECTION_{section_number}_{item_number}}}}}"])


def scrub_core_properties(document: DocumentObject, name: str) -> None:
    props = document.core_properties
    props.author = "VitaMine"
    props.last_modified_by = "VitaMine"
    props.title = name
    props.subject = "Private reusable Word CV layout"
    props.comments = ""
    props.keywords = "VitaMine Word CV template"


def clean_word_xml(data: bytes) -> bytes:
    try:
        root = etree.fromstring(data)
    except etree.XMLSyntaxError:
        return data
    namespaces = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    for node in root.xpath("//w:del|//w:commentRangeStart|//w:commentRangeEnd|//w:commentReference", namespaces=namespaces):
        parent = node.getparent()
        if parent is not None:
            parent.remove(node)
    for node in root.xpath("//w:ins", namespaces=namespaces):
        parent = node.getparent()
        if parent is None:
            continue
        position = parent.index(node)
        for child in list(node):
            parent.insert(position, child)
            position += 1
        parent.remove(node)
    for element in root.iter():
        for attribute in list(element.attrib):
            if etree.QName(attribute).localname.startswith("rsid"):
                del element.attrib[attribute]
    return etree.tostring(root, encoding="utf-8", xml_declaration=True, standalone=True)


def clean_relationships(data: bytes) -> bytes:
    try:
        root = etree.fromstring(data)
    except etree.XMLSyntaxError:
        return data
    for relationship in list(root):
        target = relationship.attrib.get("Target", "").casefold()
        rel_type = relationship.attrib.get("Type", "").casefold()
        if (
            relationship.attrib.get("TargetMode") == "External"
            or "customxml" in target
            or "customxml" in rel_type
            or "custom-properties" in rel_type
            or "comments" in target
            or "comments" in rel_type
            or "thumbnail" in target
            or "thumbnail" in rel_type
        ):
            root.remove(relationship)
    return etree.tostring(root, encoding="utf-8", xml_declaration=True)


def clean_content_types(data: bytes) -> bytes:
    try:
        root = etree.fromstring(data)
    except etree.XMLSyntaxError:
        return data
    for node in list(root):
        part = node.attrib.get("PartName", "").casefold()
        if "customxml" in part or "custom.xml" in part or "comments" in part or "thumbnail" in part:
            root.remove(node)
    return etree.tostring(root, encoding="utf-8", xml_declaration=True)


def clean_docx_package(data: bytes) -> bytes:
    source_buffer = io.BytesIO(data)
    target_buffer = io.BytesIO()
    with zipfile.ZipFile(source_buffer, "r") as source, zipfile.ZipFile(target_buffer, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            name = info.filename
            lowered = name.casefold()
            if (
                lowered.startswith("customxml/")
                or lowered == "docprops/custom.xml"
                or lowered.startswith("docprops/thumbnail.")
                or "comments" in lowered
            ):
                continue
            content = source.read(name)
            if name.endswith(".rels"):
                content = clean_relationships(content)
            elif name == "[Content_Types].xml":
                content = clean_content_types(content)
            elif name.startswith("word/") and name.endswith(".xml"):
                content = clean_word_xml(content)
            clean_info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            clean_info.compress_type = zipfile.ZIP_DEFLATED
            clean_info.create_system = info.create_system
            clean_info.external_attr = info.external_attr
            target.writestr(clean_info, content)
    return target_buffer.getvalue()


def analyze_and_skeletonize(
    data: bytes,
    name: str,
    con: sqlite3.Connection,
    *,
    llm_json: Callable[[str, dict[str, Any], dict[str, str]], tuple[dict[str, Any] | None, str | None]] | None = None,
    settings: dict[str, str] | None = None,
) -> tuple[bytes, dict[str, Any]]:
    validate_docx_bytes(data)
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:
        raise ValueError("Word could not read this DOCX document.") from exc
    units = document_units(document)
    text = "\n".join(unit.text for unit in units if unit.text)
    page_count = page_count_from_docx(data)
    headings = [unit for unit in units if probable_section_heading(unit)]
    profile, reason = deterministic_profile(text, page_count=page_count, heading_count=len(headings))
    mappings = {
        normalized_heading(unit.text): section_key
        for unit in units
        if (section_key := alias_section_key(unit.text))
    }
    method = "deterministic"
    llm_warning = ""
    configured = settings_for_llm_task(settings or {}, "custom_template_analysis")
    if llm_json is not None and str(configured.get("provider") or "none") != "none" and text:
        try:
            result, warning = llm_json(llm_prompt(units, page_count, profile), LLM_ANALYSIS_SCHEMA, configured)
        except Exception as exc:
            result, warning = None, str(exc)
        llm_warning = clean_text(warning or "")[:400]
        if isinstance(result, dict):
            proposed_profile = str(result.get("content_profile") or "")
            confidence = str(result.get("confidence") or "")
            compact_long_conflict = (
                proposed_profile == "long"
                and page_count is not None
                and page_count <= 5
                and len(re.findall(r"\b\w+\b", text)) < 2_400
            )
            if proposed_profile in CONTENT_PROFILES and confidence in {"high", "medium"} and not compact_long_conflict:
                profile = proposed_profile
                reason = f"LLM classification ({confidence} confidence)"
            exact_texts = {unit.text: normalized_heading(unit.text) for unit in units if unit.text}
            section_heading_texts = {unit.text for unit in units if unit.text and probable_section_heading(unit)}
            for item in result.get("heading_mappings") or []:
                if not isinstance(item, dict):
                    continue
                source_text = clean_text(item.get("source_text") or "")
                section_key = str(item.get("section_key") or "")
                if (
                    source_text in exact_texts
                    and source_text in section_heading_texts
                    and section_key in SECTION_ALIASES
                    and plausible_llm_heading_mapping(source_text, section_key)
                ):
                    normalized_source = exact_texts[source_text]
                    if normalized_source not in mappings:
                        mappings[normalized_source] = section_key
            method = "llm+deterministic"
    sections, mapped_sections = semantic_section_blueprint(units, mappings)
    identity_slots = masthead_identity_slots(units, mappings)
    if not any(section.get("record_unit_indices") for section in sections if not str(section["section_key"]).startswith("__unmapped_")):
        raise ValueError(
            "VitaMine could not identify reusable content sections in this Word CV. "
            "Use a DOCX with visible academic section headings and at least one entry."
        )
    person_fields = skeletonize_person(document, units, person_values(con), mappings)
    skeletonize_semantic_blueprint(units, identity_slots, sections)
    scrub_core_properties(document, name)
    buffer = io.BytesIO()
    document.save(buffer)
    skeleton = clean_docx_package(buffer.getvalue())
    validate_docx_bytes(skeleton)
    analysis = {
        "schema_version": 2,
        "content_profile": profile,
        "classification_reason": reason,
        "analysis_method": method,
        "analysis_model": str(configured.get("api_model") or configured.get("ollama_model") or "") if method.startswith("llm") else "",
        "page_count": page_count,
        "body_units": len(units),
        "table_rows": sum(unit.kind == "row" for unit in units),
        "mapped_sections": mapped_sections,
        "manual_sections": [
            clean_text(section.get("source_heading") or section.get("section_key") or "")
            for section in sections
            if str(section.get("section_key") or "").startswith("__unmapped_")
            or str(section.get("section_key") or "") in MANUAL_ONLY_SECTION_KEYS
        ],
        "heading_mappings": mappings,
        "identity_slots": identity_slots,
        "sections": sections,
        "person_fields": person_fields,
        "llm_warning": llm_warning,
    }
    return skeleton, analysis


def canonical_content(document: DocumentObject) -> dict[str, list[list[str]]]:
    content: dict[str, list[list[str]]] = defaultdict(list)
    current: str | None = None
    for unit in document_units(document):
        section_key = alias_section_key(unit.text)
        if section_key:
            current = section_key
            continue
        if current and unit.text and content_unit(unit):
            # The formal long-CV builder stores every publication as a
            # paragraph inside one single-cell table row.  A row-level join
            # turned hundreds of citations into one 50k-character paragraph.
            if unit.kind == "row" and len(unit.cell_paragraphs) == 1:
                paragraph_values = [clean_text(paragraph.text) for paragraph in unit.cell_paragraphs[0]]
                paragraph_values = [value for value in paragraph_values if value]
                if len(paragraph_values) > 1:
                    content[current].extend([[value] for value in paragraph_values])
                    continue
            content[current].append(unit.values)
    return content


def database_research_skill_items(con: sqlite3.Connection) -> list[list[str]]:
    try:
        exists = con.execute(
            "SELECT 1 FROM sqlite_master WHERE type='table' AND name='biosketch_contributions'"
        ).fetchone()
        if not exists:
            return []
        rows = con.execute(
            "SELECT title, narrative FROM biosketch_contributions ORDER BY ordinal, id LIMIT 8"
        ).fetchall()
    except sqlite3.OperationalError:
        return []
    skills: list[list[str]] = []
    for row in rows:
        title = clean_text(row["title"] if isinstance(row, sqlite3.Row) else row[0])
        narrative = clean_text(row["narrative"] if isinstance(row, sqlite3.Row) else row[1])
        value = title or re.split(r"(?<=[.!?])\s+", narrative, maxsplit=1)[0]
        if value:
            skills.append([value])
    return skills


def source_section_items(
    section_key: str, items: dict[str, list[list[str]]], con: sqlite3.Connection
) -> list[list[str]]:
    if section_key in SOURCE_SECTION_FALLBACKS:
        combined: list[list[str]] = []
        for canonical_key in SOURCE_SECTION_FALLBACKS[section_key]:
            combined.extend(items.get(canonical_key, []))
        if section_key == "research_skills" and not combined:
            combined.extend(database_research_skill_items(con))
        return combined
    return list(items.get(section_key, []))


def identity_value(role: str, values: dict[str, str]) -> str:
    if role == "display_name":
        return values.get("display_name") or values.get("full_name", "")
    if role == "position_title":
        return values.get("position_title", "")
    if role == "contact_line":
        address = values.get("office_address") or values.get("home_address") or ""
        pieces = [clean_text(address), values.get("work_phone", ""), values.get("work_email", "")]
        return ", ".join(piece for piece in pieces if piece)
    if role == "degrees":
        return values.get("degrees", "")
    if role in {"first_name", "last_name"}:
        name = values.get("full_name") or values.get("display_name", "")
        pieces = name.split()
        return " ".join(pieces[:-1]) if role == "first_name" else (pieces[-1] if pieces else "")
    if role == "institution":
        return values.get("own_institution_name") or values.get("office_address", "")
    if role == "orcid_id":
        return values.get("orcid_id", "")
    return ""


def clone_unit_after(unit: Unit) -> Unit:
    element = copy.deepcopy(unit.element)
    unit.element.addnext(element)
    if unit.kind == "paragraph":
        paragraph = Paragraph(element, unit.parent)
        return Unit(-1, "paragraph", element, unit.parent, [paragraph], [])
    row = _Row(element, unit.parent)
    cells = unique_cells(row)
    groups = [list(cell.paragraphs) for cell in cells]
    return Unit(-1, "row", element, unit.parent, [paragraph for group in groups for paragraph in group], groups)


def remove_unit(unit: Unit) -> None:
    parent = unit.element.getparent()
    if parent is not None:
        parent.remove(unit.element)


def placeholder_unit_after(heading: Unit) -> Unit:
    """Create a plain body paragraph after a section heading when needed."""

    if heading.kind != "paragraph":
        return clone_unit_after(heading)
    element = OxmlElement("w:p")
    heading.element.addnext(element)
    paragraph = Paragraph(element, heading.parent)
    return Unit(-1, "paragraph", element, heading.parent, [paragraph], [])


def ensure_section_heading_spacing(heading: Unit | None) -> None:
    """Add a modest gap when a source heading defines no preceding spacing."""

    if heading is None or heading.kind != "paragraph":
        return
    paragraph = heading.paragraphs[0]
    direct = paragraph.paragraph_format.space_before
    inherited = paragraph.style.paragraph_format.space_before if paragraph.style else None
    if direct is None and inherited is None:
        paragraph.paragraph_format.space_before = Pt(6)


def render_template(
    skeleton: bytes,
    blueprint: dict[str, Any],
    canonical_docx: Path,
    output: Path,
    con: sqlite3.Connection,
) -> dict[str, Any]:
    validate_docx_bytes(skeleton)
    try:
        document = Document(io.BytesIO(skeleton))
        canonical = Document(canonical_docx)
    except Exception as exc:
        raise ValueError("The Word template could not be rendered.") from exc
    units = document_units(document)
    by_index = {unit.index: unit for unit in units}
    items = canonical_content(canonical)
    rendered_sections: list[str] = []
    manual_sections: list[str] = []
    rendered_items = 0
    values = person_values(con)
    if int(blueprint.get("schema_version") or 1) >= 2 and blueprint.get("sections"):
        for slot in blueprint.get("identity_slots") or []:
            try:
                unit = by_index[int(slot["unit_index"])]
            except (KeyError, TypeError, ValueError):
                continue
            set_identity_slot_value(unit, slot, identity_value(str(slot.get("role") or ""), values))

        for section in blueprint.get("sections") or []:
            section_key = str(section.get("section_key") or "")
            if section_key in PRESERVED_SECTION_KEYS:
                continue
            try:
                heading = by_index[int(section["heading_unit_index"])]
            except (KeyError, TypeError, ValueError):
                heading = None
            ensure_section_heading_spacing(heading)
            body_units = [
                by_index[index]
                for raw_index in section.get("body_unit_indices") or []
                if isinstance(raw_index, int) and (index := raw_index) in by_index
            ]
            target_units = [
                by_index[index]
                for raw_index in section.get("record_unit_indices") or []
                if isinstance(raw_index, int) and (index := raw_index) in by_index
            ]
            discard_units = [
                by_index[index]
                for raw_index in section.get("discard_unit_indices") or []
                if isinstance(raw_index, int) and (index := raw_index) in by_index
            ]
            source_items = [] if section_key.startswith("__unmapped_") else source_section_items(section_key, items, con)
            manual_only = section_key.startswith("__unmapped_") or section_key in MANUAL_ONLY_SECTION_KEYS
            if manual_only or not source_items:
                placeholder_unit = target_units[0] if target_units else (
                    placeholder_unit_after(heading) if heading is not None else None
                )
                if placeholder_unit is not None:
                    set_unit_values(placeholder_unit, [MANUAL_SECTION_PLACEHOLDER])
                    for unit in body_units:
                        if unit.element is not placeholder_unit.element:
                            remove_unit(unit)
                    manual_sections.append(
                        clean_text(section.get("source_heading") or (heading.text if heading is not None else section_key))
                    )
                elif heading is not None:
                    remove_unit(heading)
                continue
            if not target_units:
                if heading is not None:
                    remove_unit(heading)
                for unit in body_units:
                    remove_unit(unit)
                continue
            while len(target_units) < len(source_items):
                target_units.append(clone_unit_after(target_units[-1]))
            for index, unit in enumerate(target_units):
                if index < len(source_items):
                    set_unit_values(unit, source_items[index])
                    rendered_items += 1
                else:
                    remove_unit(unit)
            for unit in discard_units:
                remove_unit(unit)
            rendered_sections.append(section_key)
    else:
        # Backward compatibility for templates uploaded before semantic
        # blueprint v2. Re-uploading them is recommended because their stored
        # skeleton cannot recover content that v1 misidentified as a slot.
        slots_by_section: dict[str, list[Unit]] = defaultdict(list)
        for slot in blueprint.get("slots") or []:
            try:
                unit = by_index[int(slot["unit_index"])]
            except (KeyError, TypeError, ValueError):
                continue
            section_key = str(slot.get("section_key") or "")
            if section_key:
                slots_by_section[section_key].append(unit)
        for section_key, target_units in slots_by_section.items():
            source_items = items.get(section_key, [])
            if not target_units:
                continue
            while len(target_units) < len(source_items):
                target_units.append(clone_unit_after(target_units[-1]))
            for index, unit in enumerate(target_units):
                if index < len(source_items):
                    set_unit_values(unit, source_items[index])
                    rendered_items += 1
                else:
                    remove_unit(unit)
            if source_items:
                rendered_sections.append(section_key)
    for paragraph in all_document_paragraphs(document):
        for field in blueprint.get("person_fields") or []:
            token = f"{{{{VITAMINE_PERSON_{str(field).upper()}}}}}"
            if token in paragraph.text:
                replace_paragraph_text(paragraph, token, values.get(str(field), ""))
    name = values.get("display_name") or values.get("full_name") or "VitaMine CV"
    document.core_properties.author = name
    document.core_properties.last_modified_by = "VitaMine"
    document.core_properties.title = str(blueprint.get("template_name") or output.stem)
    output.parent.mkdir(parents=True, exist_ok=True)
    buffer = io.BytesIO()
    document.save(buffer)
    output.write_bytes(clean_docx_package(buffer.getvalue()))
    return {
        "rendered_sections": rendered_sections,
        "rendered_items": rendered_items,
        "manual_sections": manual_sections,
        "unmapped_sections": sorted(set(items) - set(rendered_sections)),
    }


def template_sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()
