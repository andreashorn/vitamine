"""Database-backed, source-data-free Compact Research CV Word template."""

from __future__ import annotations

import copy
import re
import sqlite3
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from .cv_dates import UNKNOWN_CV_DATE, cv_date_sort_key, cv_end_date
from .scripts.export_publication_selection import selected_or_fallback_publications
from .scripts.export_utils import sanitize_docx_compatibility_markup


TEMPLATE_PATH = Path(__file__).parent / "static/export-templates/compact-research-cv.docx"
POSITION_SECTIONS = {
    "postdoctoral_training", "academic_appointments", "hospital_appointments",
    "professional_positions", "research_experience",
}
ACTIVITY_SECTIONS = {
    "editorial_activities", "grant_review", "committee_service",
    "professional_societies", "community_service",
}


def _text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _value(row: sqlite3.Row | None, field: str) -> str:
    return _text(row[field]) if row is not None and field in row.keys() else ""


def _detail(row: sqlite3.Row) -> str:
    # Raw import text, private source notes and person raw_json are deliberately
    # outside the field map. Pipe-separated descriptions retain their content.
    values = [_value(row, key) for key in ("title", "organization", "location", "role", "amount")]
    values.extend(_value(row, "description").split("|"))
    result: list[str] = []
    seen: set[str] = set()
    for value in values:
        value = _text(value)
        key = re.sub(r"\W+", " ", value.casefold()).strip()
        if value and key not in seen:
            result.append(value)
            seen.add(key)
    return "; ".join(result)


def _period(row: sqlite3.Row) -> str:
    start, end = _value(row, "start_date"), _value(row, "end_date")
    if start and end and start != end:
        return f"{start}–{end}"
    return start or end


def _active_funding(row: sqlite3.Row, today: date) -> bool:
    status = _value(row, "grant_status").casefold()
    if not status:
        status = "submitted" if _value(row, "subcategory").casefold() == "grant_application" else "funded"
    if status != "funded":
        return False
    start = cv_date_sort_key(_value(row, "start_date"))
    if start != UNKNOWN_CV_DATE and start > (today.year, today.month, today.day):
        return False
    try:
        end = cv_end_date(_value(row, "end_date"))
    except ValueError:
        return False
    return end is None or end >= today


def _citation(row: sqlite3.Row) -> str:
    # A complete stored citation preserves volume, issue and pages, which do
    # not have dedicated database columns. Otherwise use structured fields.
    citation = _value(row, "raw_citation")
    # A stored citation is reusable only while its structured bibliographic
    # fields still agree. Manual database edits must appear in the export.
    normalize = lambda value: re.sub(r"\W+", " ", value.casefold()).strip()
    if citation and any(
        value and normalize(value) not in normalize(citation)
        for key in ("authors", "title", "venue", "year")
        if (value := _value(row, key))
    ):
        citation = ""
    if not citation:
        citation = ". ".join(
            value.rstrip(".") for key in ("authors", "title", "venue", "year")
            if (value := _value(row, key))
        )
    doi = re.sub(r"^(?:https?://(?:dx\.)?doi\.org/|doi:\s*)", "", _value(row, "doi"), flags=re.I)
    if doi and doi.casefold() not in citation.casefold():
        citation = f"{citation.rstrip('. ')}. https://doi.org/{doi}".lstrip(". ")
    return citation


def _name_pattern(person: sqlite3.Row | None) -> re.Pattern[str]:
    variants = []
    for key in ("display_name", "full_name"):
        name = _value(person, key)
        if not name:
            continue
        variants.append(re.escape(name))
        if "," in name:
            surname, given = (part.strip() for part in name.split(",", 1))
        else:
            given, _, surname = name.rpartition(" ")
        if given and surname:
            variants.append(rf"{re.escape(surname)},?\s+{re.escape(given[0])}\.?")
    return re.compile(r"(?<!\w)(?:" + "|".join(variants) + r")(?!\w)" if variants else r"(?!x)x", re.I)


def _emphasize_name(paragraph: Paragraph, pattern: re.Pattern[str]) -> None:
    value = paragraph.text
    matches = list(pattern.finditer(value))
    if not matches:
        return
    properties = copy.deepcopy(paragraph.runs[0]._r.rPr)
    paragraph.clear()
    cursor = 0
    for match in matches:
        for text, bold in ((value[cursor:match.start()], False), (match.group(), True)):
            if text:
                run = paragraph.add_run(text)
                if properties is not None:
                    run._r.insert(0, copy.deepcopy(properties))
                run.bold = bold
        cursor = match.end()
    if cursor < len(value):
        run = paragraph.add_run(value[cursor:])
        if properties is not None:
            run._r.insert(0, copy.deepcopy(properties))


def _replace(paragraph: Paragraph, value: str) -> None:
    properties = copy.deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs else None
    paragraph.clear()
    run = paragraph.add_run(value)
    if properties is not None:
        run._r.insert(0, properties)


def _remove(element: Any) -> None:
    element.getparent().remove(element)


def render_compact_research_cv(
    con: sqlite3.Connection, output_path: Path, *, today: date | None = None,
) -> dict[str, Any]:
    """Render the Short CV selections; omit missing data without placeholders.

    Word PAGE/NUMPAGES fields update on open and on PDF export. There is no
    fixed page cap: long entries flow without truncation or smaller typography.
    """
    today = today or date.today()
    document = Document(TEMPLATE_PATH)
    body = list(document.element.body)
    person = con.execute("SELECT * FROM person WHERE id=1").fetchone()
    identity = [
        _value(person, "display_name") or _value(person, "full_name"),
        _value(person, "degrees"), _value(person, "orcid_id"),
        _value(person, "position_title"),
        _value(person, "office_address") or _value(person, "own_institution_name"),
    ]
    personal = Table(body[2], document)
    for row, value in zip(list(personal.rows), identity):
        if value:
            _replace(row.cells[1].paragraphs[0], value)
        else:
            _remove(row._tr)
    if not any(identity):
        _remove(body[1]); _remove(body[2])

    entries = con.execute("SELECT * FROM cv_entries WHERE include_short=1 ORDER BY id").fetchall()
    # Unknown dates follow dated records; source values remain unchanged.
    entries.sort(key=lambda row: (
        cv_date_sort_key(_value(row, "start_date") or _value(row, "end_date")) != UNKNOWN_CV_DATE,
        cv_date_sort_key(_value(row, "start_date") or _value(row, "end_date")), row["id"],
    ), reverse=True)
    groups: dict[str, list[tuple[str, str]]] = {
        key: [] for key in ("education", "positions", "honors", "activities", "funding")
    }
    seen: set[tuple[str, str, str]] = set()
    for row in entries:
        section = row["section_key"]
        key = (
            "positions" if section in POSITION_SECTIONS else
            "activities" if section in ACTIVITY_SECTIONS else section
        )
        if key not in groups or (key == "funding" and not _active_funding(row, today)):
            continue
        detail, period = _detail(row), _period(row)
        marker = (key, period.casefold(), detail.casefold())
        if detail and marker not in seen:
            groups[key].append((period, detail))
            seen.add(marker)

    publications = selected_or_fallback_publications(con, profile="short", limit=5)
    citations = [citation for row in publications if (citation := _citation(row))]
    name_pattern = _name_pattern(person)
    counts: dict[str, int] = {}
    for key, heading_index, table_index in (("education", 3, 4), ("positions", 5, 6), ("honors", 7, 8)):
        records = groups[key]
        counts[key] = len(records)
        if not records:
            _remove(body[heading_index]); _remove(body[table_index])
            continue
        table = Table(body[table_index], document)
        prototype = copy.deepcopy(table.rows[0]._tr)
        _remove(table.rows[0]._tr)
        for period, detail in records:
            table._tbl.append(copy.deepcopy(prototype))
            for cell, value in zip(table.rows[-1].cells, (period, detail)):
                _replace(cell.paragraphs[0], value)

    for key, heading_index, paragraph_index in (("activities", 9, 10), ("funding", 11, 12), ("publications", 13, 14)):
        values = citations if key == "publications" else [" ".join(filter(None, record)) for record in groups[key]]
        counts[key] = len(values)
        if not values:
            _remove(body[heading_index]); _remove(body[paragraph_index])
            continue
        if key == "publications" and len(values) < 5:
            _replace(Paragraph(body[heading_index], document), "■ Key papers")
        prototype = body[paragraph_index]
        for value in values:
            element = copy.deepcopy(prototype)
            prototype.addprevious(element)
            _replace(Paragraph(element, document), value)
            if key == "publications":
                _emphasize_name(Paragraph(element, document), name_pattern)
        _remove(prototype)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    sanitize_docx_compatibility_markup(output_path)
    return {
        "rendered_sections": [key for key, count in counts.items() if count],
        "rendered_items": sum(counts.values()), "section_counts": counts,
        "manual_sections": [], "publication_limit": 5,
        "selection_profile": "short", "page_numbering": "PAGE/NUMPAGES",
    }
